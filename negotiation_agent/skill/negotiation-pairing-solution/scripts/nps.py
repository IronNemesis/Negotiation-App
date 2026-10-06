"""Negotiation Pairing Solution command-line tool.

Every command prints one JSON object to stdout. On failure the object has
"ok": false, an "error" code, and a plain-language "message". Run
`python nps.py <command> --help` for the options of each command.

Commands (task IDs refer to negotiation_agent/agent/task-specs):
  snapshot        T1  read roster, pairing history, and simulation settings
  match-absences  T2  match the instructor's absence list to roster students
  pair            T3+T4  generate a pairing draft and validate it
  validate        T4  re-check Pairing Draft.xlsx after the instructor edits it
  record          T5  write an approved pairing into Pairing History
  mailmerge       T6  build Role Mail Merge.xlsx for the Outlook mail merge
  survey          T7+T8  read the Qualtrics export and match it to students
  outcomes        T9  build Outcomes.xlsx in pairing order
  log             write an instructor decision to the Run Log
"""
from __future__ import annotations

import argparse
import csv
import difflib
import json
import random
import re
import shutil
import sys
import time
from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime
from itertools import combinations
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    print(json.dumps({"ok": False, "error": "missing_dependency",
                      "message": "The Python package openpyxl is not installed. Install it with: pip install openpyxl"}))
    sys.exit(1)

PAIRING_SHEET, CHECKS_SHEET, INFO_SHEET = "Pairing", "Checks", "Info"
DRAFT_HEADERS = ["Pair #", "Student ID", "Student Name", "Role", "Partner", "Repeat Partner?"]
MAIL_MERGE_HEADERS = ["First Name", "Last Name", "Email", "Simulation", "Role", "Role Information", "Survey Link", "Status"]
OBSERVER, INSTRUCTOR = "Observer", "Instructor"
HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="D9E7DF")


class NPSError(Exception):
    def __init__(self, code: str, message: str, details: dict | None = None):
        super().__init__(message)
        self.code, self.message, self.details = code, message, details or {}


# ---------------------------------------------------------------- helpers

def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def low(value) -> str:
    return clean(value).lower()


def norm_id(value) -> str:
    """Compare student IDs the way a person would: ignore spaces and leading zeros."""
    s = re.sub(r"\s+", "", clean(value))
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s.lstrip("0") or s


def find_header(rows, required, limit=15):
    for i, row in enumerate(rows[:limit]):
        cells = [low(c) for c in row]
        if all(any(word in c for c in cells) for word in required):
            return i, cells
    return None, None


def find_col(cells, *words, exclude=()):
    for i, c in enumerate(cells):
        if all(w in c for w in words) and not any(x in c for x in exclude):
            return i
    return None


def cell(row, index):
    return row[index] if index is not None and index < len(row) else None


def read_rows(path: Path):
    if not path.exists():
        raise NPSError("file_missing", f"Could not find {path.name} in {path.parent}.", {"path": str(path)})
    try:
        return list(load_workbook(path, data_only=True).worksheets[0].iter_rows(values_only=True))
    except PermissionError:
        raise NPSError("file_locked", f"{path.name} is open in Excel or locked by OneDrive. Close it and try again.",
                       {"path": str(path)})
    except Exception as exc:  # corrupt or not an Excel file
        raise NPSError("file_unreadable", f"Could not read {path.name}: {exc}", {"path": str(path)})


def save_workbook(wb, path: Path):
    try:
        wb.save(path)
    except PermissionError:
        raise NPSError("file_locked", f"{path.name} is open in Excel or locked by OneDrive. Close it and try again.",
                       {"path": str(path)})


def style_header_row(ws, row=1):
    for c in ws[row]:
        if c.value is not None:
            c.font, c.fill = HEADER_FONT, HEADER_FILL


def stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H%M%S")


def parse_number(value):
    s = clean(value).replace("$", "").replace(",", "")
    if s == "":
        return None
    try:
        n = float(s)
        return int(n) if n.is_integer() else n
    except ValueError:
        return clean(value)


# ---------------------------------------------------------------- folder and file loading

class Paths:
    def __init__(self, root: str, section: str, run: str | None = None):
        self.root = Path(root)
        if not self.root.is_dir():
            raise NPSError("root_missing", f"The course folder {self.root} does not exist.")
        self.section_dir = self.root / section
        if not self.section_dir.is_dir():
            available = sorted(p.name for p in self.root.iterdir() if p.is_dir() and p.name != "Templates")
            raise NPSError("section_missing", f"There is no section folder named '{section}'.", {"available": available})
        self.section = section
        self.settings = self.root / "Simulation Settings.xlsx"
        self.templates = self.root / "Templates"
        self.roster = self.section_dir / "Roster.xlsx"
        self.history = self.section_dir / "Pairing History.xlsx"
        self.run_log = self.section_dir / "Run Log.xlsx"
        self.backups = self.section_dir / "Backups"
        self.run = run

    @property
    def run_dir(self):
        return self.section_dir / self.run if self.run else None

    def require_run(self):
        if not self.run_dir or not self.run_dir.is_dir():
            runs = sorted(p.name for p in self.section_dir.iterdir() if p.is_dir() and p.name != "Backups")
            raise NPSError("run_missing", f"There is no simulation folder named '{self.run}' in {self.section}.",
                           {"available": runs})
        return self.run_dir

    def backup(self, path: Path, label: str) -> str:
        self.backups.mkdir(exist_ok=True)
        target = self.backups / f"{path.stem} {stamp()} {label}{path.suffix}"
        shutil.copy2(path, target)
        return target.name


def load_settings(paths: Paths, simulation: str) -> dict:
    rows = read_rows(paths.settings)
    i, cells = find_header(rows, ["simulation", "role 1"])
    if i is None:
        raise NPSError("settings_layout", "Simulation Settings.xlsx needs a header row with 'Simulation' and 'Role 1' columns.")
    c = {key: find_col(cells, *words, exclude=exclude) for key, words, exclude in [
        ("sim", ("simulation",), ()), ("r1", ("role 1",), ("info",)), ("r2", ("role 2",), ("info",)),
        ("i1", ("role 1", "info"), ()), ("i2", ("role 2", "info"), ()), ("link", ("survey link",), ()),
        ("sid", ("student id",), ()), ("map", ("mapping",), ()), ("tpl", ("template",), ()),
        ("bal", ("balanc",), ())]}
    names = []
    for r in rows[i + 1:]:
        name = clean(cell(r, c["sim"]))
        if not name:
            continue
        names.append(name)
        if name.lower() != simulation.strip().lower():
            continue
        roles = [clean(cell(r, c["r1"])), clean(cell(r, c["r2"]))]
        if not all(roles):
            raise NPSError("settings_roles", f"{name} needs both Role 1 and Role 2 filled in Simulation Settings.")
        mapping = []
        for part in clean(cell(r, c["map"])).split(";"):
            if "=" in part:
                src, dst = part.split("=", 1)
                mapping.append((src.strip(), dst.strip()))
        return {
            "simulation": name, "roles": roles,
            "role_info": {roles[0]: clean(cell(r, c["i1"])), roles[1]: clean(cell(r, c["i2"]))},
            "survey_link": clean(cell(r, c["link"])), "survey_id_column": clean(cell(r, c["sid"])) or "Q1",
            "mapping": mapping, "template": clean(cell(r, c["tpl"])),
            "balancing": low(cell(r, c["bal"])) in ("yes", "y", "true", "on", "1"),
        }
    raise NPSError("simulation_missing", f"'{simulation}' is not listed in Simulation Settings.xlsx.", {"available": names})


def split_name(raw: str, first: str = "", last: str = ""):
    if first or last:
        return first, last
    if "," in raw:
        last, first = [p.strip() for p in raw.split(",", 1)]
        return first, last
    parts = raw.split()
    return (" ".join(parts[:-1]), parts[-1]) if len(parts) > 1 else (raw, "")


def load_roster(paths: Paths) -> dict:
    rows = read_rows(paths.roster)
    i, cells = find_header(rows, ["student id"])
    if i is None:
        i, cells = find_header(rows, ["emplid"])
    if i is None:
        raise NPSError("roster_layout", "Roster.xlsx has no 'Student ID' column header.")
    c_id = find_col(cells, "student id") if find_col(cells, "student id") is not None else find_col(cells, "emplid")
    c_name = find_col(cells, "name", exclude=("first", "last", "preferred"))
    c_first, c_last = find_col(cells, "first"), find_col(cells, "last")
    c_email = find_col(cells, "email")
    c_status = next((j for j, h in enumerate(cells) if h.startswith("status") and "note" not in h), None)
    c_note = find_col(cells, "note")
    if c_name is None and (c_first is None or c_last is None):
        raise NPSError("roster_layout", "Roster.xlsx needs a 'Name' column, or 'First Name' and 'Last Name' columns.")
    pairable, excluded = [], []
    for rownum, r in enumerate(rows[i + 1:], start=i + 2):
        if all(v is None for v in r):
            continue
        first, last = split_name(clean(cell(r, c_name)), clean(cell(r, c_first)), clean(cell(r, c_last)))
        sid = clean(cell(r, c_id))
        status = clean(cell(r, c_status)) if c_status is not None else "Enrolled"
        rec = {"id": sid, "key": norm_id(sid), "first": first, "last": last, "name": f"{first} {last}".strip(),
               "email": clean(cell(r, c_email)), "status": status, "note": clean(cell(r, c_note)), "row": rownum}
        if not sid:
            excluded.append({**rec, "reason": "No student ID"})
        elif status.lower() != "enrolled":
            excluded.append({**rec, "reason": f"Status is '{status}'"})
        else:
            pairable.append(rec)
    dupes = [k for k, n in Counter(s["key"] for s in pairable).items() if n > 1]
    if dupes:
        names = [f"{s['name']} ({s['id']}, row {s['row']})" for s in pairable if s["key"] in dupes]
        raise NPSError("duplicate_ids", "The roster has the same student ID more than once.", {"students": names})
    return {"pairable": pairable, "excluded": excluded, "by_key": {s["key"]: s for s in pairable + excluded if s["key"]}}


def load_history(paths: Paths) -> dict:
    rows = read_rows(paths.history)
    i, cells = find_header(rows, ["student id"])
    if i is None:
        raise NPSError("history_layout", "Pairing History.xlsx has no 'Student ID' column header.")
    c_id, c_name = find_col(cells, "student id"), find_col(cells, "name")
    rounds = []
    for j, h in enumerate(cells):
        if "pair" in h and j not in (c_id, c_name):
            role_col = j + 1 if j + 1 < len(cells) and "role" in cells[j + 1] else None
            rounds.append({"pair_col": j, "role_col": role_col, "label": clean(rows[i][j])})
    entries = []
    for rownum, r in enumerate(rows[i + 1:], start=i + 2):
        sid = clean(cell(r, c_id))
        if not sid and not clean(cell(r, c_name)):
            continue
        entries.append({
            "row": rownum, "id": sid, "key": norm_id(sid), "name": clean(cell(r, c_name)),
            "ids": {rd["pair_col"]: clean(cell(r, rd["pair_col"])) for rd in rounds if clean(cell(r, rd["pair_col"]))},
            "roles": Counter(low(cell(r, rd["role_col"])) for rd in rounds
                             if rd["role_col"] is not None and clean(cell(r, rd["role_col"]))),
        })
    shared = defaultdict(list)  # frozenset(key_a, key_b) -> pair IDs they already share
    for rd in rounds:
        groups = defaultdict(list)
        for e in entries:
            if rd["pair_col"] in e["ids"] and e["key"]:
                groups[e["ids"][rd["pair_col"]]].append(e["key"])
        for pid, keys in groups.items():
            for a, b in combinations(sorted(set(keys)), 2):
                shared[frozenset((a, b))].append(pid)
    all_ids = [v for e in entries for v in e["ids"].values()]
    return {"header_row": i + 1, "cells": cells, "rounds": rounds, "entries": entries,
            "by_key": {e["key"]: e for e in entries if e["key"]}, "shared": shared,
            "format": detect_pair_id_format(all_ids), "tracks_roles": any(rd["role_col"] is not None for rd in rounds)}


def detect_pair_id_format(values) -> dict:
    values = sorted(set(values))
    if not values:
        return {"kind": "round", "prefix": "N", "sep": "-", "width": 2, "next": 1}
    matches = [re.fullmatch(r"([A-Za-z]*)(\d+)([-_.])(\d+)", v) for v in values]
    if all(matches) and len({(m.group(1), m.group(3)) for m in matches}) == 1:
        m0 = matches[0]
        return {"kind": "round", "prefix": m0.group(1), "sep": m0.group(3),
                "width": max(len(m.group(4)) for m in matches), "next": max(int(m.group(2)) for m in matches) + 1}
    matches = [re.fullmatch(r"([A-Za-z _-]*?)(\d+)", v) for v in values]
    if all(matches) and len({m.group(1) for m in matches}) == 1:
        return {"kind": "counter", "prefix": matches[0].group(1), "width": max(len(m.group(2)) for m in matches),
                "next": max(int(m.group(2)) for m in matches) + 1}
    raise NPSError("pair_id_format_unknown", "The pair IDs in Pairing History do not follow one recognizable pattern.",
                   {"examples": values[:10]})


def new_pair_ids(fmt: dict, count: int):
    if fmt["kind"] == "round":
        return [f"{fmt['prefix']}{fmt['next']}{fmt['sep']}{n:0{fmt['width']}d}" for n in range(1, count + 1)]
    return [f"{fmt['prefix']}{n:0{fmt['width']}d}" for n in range(fmt["next"], fmt["next"] + count)]


def history_link_report(roster, history):
    unlinked = [e["name"] or e["id"] for e in history["entries"] if e["key"] not in roster["by_key"]]
    not_enrolled = [s["name"] for s in roster["excluded"] if s["key"] in history["by_key"]]
    return unlinked, not_enrolled


# ---------------------------------------------------------------- run log

def write_log(paths: Paths, step: str, details: str) -> str | None:
    """Append one plain-language line to the section's Run Log. Returns a warning if it could not be written."""
    try:
        if paths.run_log.exists():
            wb = load_workbook(paths.run_log)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Run Log"
            ws.append(["Time", "Simulation Folder", "Step", "Details"])
            style_header_row(ws)
            for col, width in zip("ABCD", (20, 28, 34, 100)):
                ws.column_dimensions[col].width = width
        ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), paths.run or "", step, details])
        wb.save(paths.run_log)
        return None
    except PermissionError:
        return f"Run Log.xlsx is open in Excel, so this entry was not logged: {step}: {details}"


def finish(result: dict, paths: Paths | None = None, step: str | None = None, details: str | None = None):
    if paths and step:
        warning = write_log(paths, step, details or "")
        if warning:
            result.setdefault("log_warning", warning)
    print(json.dumps(result, indent=2, ensure_ascii=False))


# ---------------------------------------------------------------- T1 snapshot

def cmd_snapshot(a):
    paths = Paths(a.root, a.section)
    settings = load_settings(paths, a.simulation)
    roster, history = load_roster(paths), load_history(paths)
    unlinked, not_enrolled = history_link_report(roster, history)
    fmt = history["format"]
    result = {
        "ok": True, "section": a.section, "simulation": settings["simulation"], "roles": settings["roles"],
        "role_balancing": settings["balancing"],
        "pairable_count": len(roster["pairable"]),
        "pairable": [{"id": s["id"], "name": s["name"], **({"note": s["note"]} if s["note"] else {})}
                     for s in roster["pairable"]],
        "excluded": [{"name": s["name"], "id": s["id"], "reason": s["reason"]} for s in roster["excluded"]],
        "missing_email": [s["name"] for s in roster["pairable"] if not s["email"]],
        "history": {"negotiations_recorded": [rd["label"] for rd in history["rounds"]],
                    "in_history_not_on_roster": unlinked, "in_history_not_enrolled": not_enrolled,
                    "next_pair_id": new_pair_ids(fmt, 1)[0], "tracks_roles": history["tracks_roles"]},
    }
    finish(result, paths, "T1 Snapshot",
           f"{settings['simulation']}: {len(roster['pairable'])} pairable, {len(roster['excluded'])} excluded, "
           f"{len(history['rounds'])} past negotiations in Pairing History.")


# ---------------------------------------------------------------- T2 absences

def simple(text: str) -> str:
    return " ".join(re.sub(r"[,.]", " ", text.lower()).split())


def match_absence(entry: str, roster: dict) -> dict:
    everyone = roster["pairable"] + [s for s in roster["excluded"] if s["key"]]
    e, ek = simple(entry), norm_id(entry)
    rules = [
        ("student ID", lambda s: any(ch.isdigit() for ch in entry) and s["key"] == ek),
        ("email", lambda s: s["email"] and s["email"].lower() == entry.strip().lower()),
        ("full name", lambda s: e in (simple(f"{s['first']} {s['last']}"), simple(f"{s['last']} {s['first']}"))),
    ]
    for rule, test in rules:
        hits = [s for s in everyone if test(s)]
        if len(hits) == 1:
            s = hits[0]
            if s in roster["excluded"]:
                return {"entry": entry, "status": "already_excluded", "name": s["name"], "id": s["id"],
                        "reason": s["reason"]}
            return {"entry": entry, "status": "matched", "id": s["id"], "name": s["name"], "rule": rule}
        if len(hits) > 1:
            return {"entry": entry, "status": "unmatched", "reason": f"Matches more than one student by {rule}",
                    "suggestions": [{"id": s["id"], "name": s["name"]} for s in hits]}
    scored = []
    for s in roster["pairable"]:
        score = max(difflib.SequenceMatcher(None, e, simple(x)).ratio()
                    for x in (s["name"], s["first"], s["last"], s["email"].split("@")[0]) if x)
        # A name token that appears exactly (a first name, a last name, or a "Goes by Katie" note) is a strong hint.
        token_hit = any(tok in simple(s["note"]).split() or tok in (simple(s["first"]), simple(s["last"]))
                        for tok in e.split())
        scored.append((score + (0.5 if token_hit else 0), token_hit, s))
    scored.sort(key=lambda x: -x[0])
    if any(hit for _, hit, _ in scored):
        scored = [x for x in scored if x[1]]
    hits = [s for score, _, s in scored[:3] if score >= 0.5]
    reason = "Matches more than one student" if len([x for x in scored if x[1]]) > 1 else "No exact match"
    return {"entry": entry, "status": "unmatched", "reason": reason,
            "suggestions": [{"id": s["id"], "name": s["name"], **({"note": s["note"]} if s["note"] else {})}
                            for s in hits]}


def cmd_match_absences(a):
    paths = Paths(a.root, a.section)
    roster = load_roster(paths)
    results = [match_absence(x, roster) for x in a.absent if x.strip()]
    matched = [r for r in results if r["status"] == "matched"]
    unmatched = [r for r in results if r["status"] == "unmatched"]
    already = [r for r in results if r["status"] == "already_excluded"]
    finish({"ok": True, "matched": matched, "unmatched": unmatched, "already_excluded": already,
            "all_matched": not unmatched},
           paths, "T2 Match absences",
           f"Entries: {[r['entry'] for r in results]}. Matched: {[m['name'] for m in matched]}. "
           f"Unmatched: {[u['entry'] for u in unmatched]}.")


# ---------------------------------------------------------------- T3 pairing

REPEAT_COST, APART_COST, ROLE_COST = 100, 10000, 1


def group_cost(group, shared, apart, leans):
    """Repeat partners cost far more than a role imbalance, so roles never force a repeat."""
    cost = 0
    for a, b in combinations(group, 2):
        pair = frozenset((a, b))
        cost += REPEAT_COST * len(shared.get(pair, ())) + (APART_COST if pair in apart else 0)
    if len(group) == 2 and leans:
        la, lb = leans.get(group[0], 0), leans.get(group[1], 0)
        if la * lb > 0:  # both lean toward the same role, so one of them has to repeat it
            cost += ROLE_COST * min(abs(la), abs(lb))
    return cost


def search_groups(keys, sizes, shared, apart, rng, leans=None, seconds=20.0):
    best, best_cost = None, None
    deadline = time.time() + seconds
    while True:
        order = keys[:]
        rng.shuffle(order)
        groups, start = [], 0
        for size in sizes:
            groups.append(order[start:start + size])
            start += size
        costs = [group_cost(g, shared, apart, leans) for g in groups]
        improved = True
        while improved and time.time() < deadline:
            improved = False
            for g1, g2 in combinations(range(len(groups)), 2):
                if costs[g1] == 0 and costs[g2] == 0:
                    continue
                for x in range(len(groups[g1])):
                    for y in range(len(groups[g2])):
                        a_, b_ = groups[g1][:], groups[g2][:]
                        a_[x], b_[y] = groups[g2][y], groups[g1][x]
                        c1, c2 = group_cost(a_, shared, apart, leans), group_cost(b_, shared, apart, leans)
                        if c1 + c2 < costs[g1] + costs[g2]:
                            groups[g1], groups[g2], costs[g1], costs[g2] = a_, b_, c1, c2
                            improved = True
        total = sum(costs)
        if best_cost is None or total < best_cost:
            best, best_cost = [g[:] for g in groups], total
        if best_cost == 0 or time.time() >= deadline:
            return best, best_cost


def role_lean(roles_count: Counter, r1: str, r2: str) -> int:
    """Positive when the student has played role 1 more often than role 2."""
    return roles_count.get(r1.lower(), 0) - roles_count.get(r2.lower(), 0)


def assign_roles(group, roles, past, balancing, rng, doubled=None):
    r1, r2 = roles
    if len(group) == 1:
        lean = role_lean(past.get(group[0], Counter()), r1, r2) if balancing else 0
        return {group[0]: r2 if lean > 0 else r1 if lean < 0 else rng.choice(roles)}
    if len(group) == 3:
        single = r2 if doubled.lower() == r1.lower() else r1
        double = r1 if single == r2 else r2
        ranked = sorted(group, key=lambda k: (past.get(k, Counter()).get(single.lower(), 0), rng.random())
                        if balancing else rng.random())
        return {k: single if k == ranked[0] else double for k in group}
    a, b = group
    if balancing:
        la, lb = role_lean(past.get(a, Counter()), r1, r2), role_lean(past.get(b, Counter()), r1, r2)
        if la != lb:
            return {a: r1, b: r2} if la < lb else {a: r2, b: r1}
    return {a: r1, b: r2} if rng.random() < 0.5 else {a: r2, b: r1}


def pair_sort_key(pair: str):
    try:
        return (0, float(pair), "")
    except ValueError:
        return (1, 0.0, pair)


def run_folder_name(simulation: str, when: str) -> str:
    return f"{simulation} {when}"


def read_info(run_dir: Path) -> dict:
    path = run_dir / "Pairing Draft.xlsx"
    if not path.exists():
        raise NPSError("draft_missing", f"There is no Pairing Draft.xlsx in {run_dir.name}. Run Stage 1 first.")
    try:
        wb = load_workbook(path)
    except PermissionError:
        raise NPSError("file_locked", "Pairing Draft.xlsx is open in Excel or locked by OneDrive. Close it and try again.")
    if INFO_SHEET not in wb.sheetnames:
        raise NPSError("draft_layout", "Pairing Draft.xlsx is missing its Info sheet. Generate a new draft.")
    info = {clean(r[0]): clean(r[1]) for r in wb[INFO_SHEET].iter_rows(values_only=True) if r and r[0]}
    return info


def write_info(ws, info: dict):
    ws.delete_rows(1, ws.max_row)
    for k, v in info.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 70
    for c in ws["A"]:
        c.font = HEADER_FONT


def cmd_pair(a):
    paths = Paths(a.root, a.section)
    settings = load_settings(paths, a.simulation)
    roster, history = load_roster(paths), load_history(paths)
    by_key = {s["key"]: s for s in roster["pairable"]}
    absent = []
    for raw in a.absent_id:
        k = norm_id(raw)
        if k not in by_key:
            raise NPSError("unknown_absent_id", f"Student ID {raw} is not a pairable student in this section.")
        absent.append(k)
    attending = [s["key"] for s in roster["pairable"] if s["key"] not in absent]
    if len(attending) < 2:
        raise NPSError("too_few_students", "Fewer than two students are attending, so there is nobody to pair.")

    odd_student = None
    if len(attending) % 2:
        if not a.odd:
            counts = {k: len(history["by_key"].get(k, {}).get("ids", {})) for k in attending}
            fewest = min(attending, key=lambda k: (counts[k], by_key[k]["name"]))
            raise NPSError("odd_count", f"{len(attending)} students are attending, which is an odd number.", {
                "attending_count": len(attending),
                "options": ["trio: one group of three (choose which role is doubled)",
                            "instructor: the instructor partners with one student",
                            "observe: one student observes this round"],
                "suggested_student": {"id": by_key[fewest]["id"], "name": by_key[fewest]["name"],
                                      "past_negotiations": counts[fewest]}})
        if a.odd in ("instructor", "observe"):
            if not a.odd_student_id or norm_id(a.odd_student_id) not in attending:
                raise NPSError("odd_student_missing", "Give --odd-student-id for an attending student.")
            odd_student = norm_id(a.odd_student_id)
        elif a.odd == "trio":
            if not a.trio_role or a.trio_role.lower() not in [r.lower() for r in settings["roles"]]:
                raise NPSError("trio_role_missing", f"Give --trio-role as one of {settings['roles']}.")
    elif a.odd:
        a.odd = None  # even count: nothing to arrange

    apart = set()
    for spec in a.keep_apart:
        ids = [norm_id(x) for x in spec.split(",")]
        if len(ids) != 2 or not all(i in by_key for i in ids):
            raise NPSError("keep_apart_invalid", f"--keep-apart '{spec}' must be two roster student IDs separated by a comma.")
        apart.add(frozenset(ids))

    pool = [k for k in attending if k != odd_student]
    sizes = [2] * (len(pool) // 2)
    if len(pool) % 2:
        sizes[-1] = 3
    seed = a.seed if a.seed is not None else random.SystemRandom().randint(1, 10**6)
    rng = random.Random(seed)
    past = {k: e["roles"] for k, e in history["by_key"].items()}
    leans = ({k: role_lean(past.get(k, Counter()), *settings["roles"]) for k in pool}
             if settings["balancing"] else None)
    groups, _ = search_groups(pool, sizes, history["shared"], apart, rng, leans)
    groups.sort(key=len)  # trio last
    repeats = sum(len(history["shared"].get(frozenset(p), ())) for g in groups for p in combinations(g, 2))
    assignments = []  # (pair number, key, role)
    for n, g in enumerate(groups, start=1):
        roles = assign_roles(g, settings["roles"], past, settings["balancing"], rng, a.trio_role)
        assignments += [(n, k, roles[k]) for k in g]
    if odd_student and a.odd == "instructor":
        role = assign_roles([odd_student], settings["roles"], past, settings["balancing"], rng)[odd_student]
        assignments.append((len(groups) + 1, odd_student, role))
    if odd_student and a.odd == "observe":
        assignments.append(("", odd_student, OBSERVER))

    when = a.date or date.today().isoformat()
    run = run_folder_name(settings["simulation"], when)
    run_dir = paths.section_dir / run
    run_dir.mkdir(exist_ok=True)
    draft_path = run_dir / "Pairing Draft.xlsx"
    replaced = None
    if draft_path.exists():
        if read_info(run_dir).get("Status", "").startswith("Recorded"):
            raise NPSError("already_recorded", f"The pairing in {run} has already been recorded in Pairing History. "
                                               "Use a different --date for a new negotiation.")
        paths.run = run
        replaced = paths.backup(draft_path, "replaced")

    wb = Workbook()
    ws = wb.active
    ws.title = PAIRING_SHEET
    ws.append(DRAFT_HEADERS)
    for n, k, role in assignments:
        s = by_key[k]
        ws.append([n, s["id"], s["name"], role, "", ""])
        ws.cell(ws.max_row, 2).number_format = "@"
    for col, width in zip("ABCDEF", (8, 13, 24, 12, 30, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    style_header_row(ws)
    wb.create_sheet(CHECKS_SHEET)
    info = {
        "Section": a.section, "Simulation": settings["simulation"], "Negotiation Date": when,
        "Absent IDs": ", ".join(by_key[k]["id"] for k in absent),
        "Absent Names": ", ".join(by_key[k]["name"] for k in absent),
        "Odd Arrangement": a.odd or "", "Odd Student ID": by_key[odd_student]["id"] if odd_student else "",
        "Doubled Role": (a.trio_role or "") if a.odd == "trio" else "",
        "Keep Apart": "; ".join(",".join(by_key[k]["id"] for k in sorted(p)) for p in apart),
        "Seed": seed, "Status": "Draft", "Created": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }
    write_info(wb.create_sheet(INFO_SHEET), info)
    save_workbook(wb, draft_path)

    paths.run = run
    checks = validate_draft(paths, roster, history, settings)
    result = {"ok": True, "run": run, "draft": str(draft_path), "seed": seed,
              "students_paired": len([x for x in assignments if x[2] != OBSERVER]),
              "groups": len(groups) + (1 if odd_student and a.odd == "instructor" else 0),
              "repeat_partners": repeats, **checks}
    if replaced:
        result["previous_draft_backup"] = replaced
    finish(result, paths, "T3 Generate pairing",
           f"Draft with seed {seed}: {result['students_paired']} students, {len(checks['errors'])} errors, "
           f"{len(checks['warnings'])} warnings. Absent: {info['Absent Names'] or 'none'}. "
           f"Odd arrangement: {a.odd or 'none'}. Keep apart: {info['Keep Apart'] or 'none'}.")


# ---------------------------------------------------------------- T4 validation

def read_draft_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    i, cells = find_header(rows, ["student id", "role"])
    if i is None:
        raise NPSError("draft_layout", "The Pairing sheet is missing its 'Student ID' or 'Role' column header.")
    cols = {"pair": find_col(cells, "pair", exclude=("partner",)), "id": find_col(cells, "student id"),
            "name": find_col(cells, "name"), "role": find_col(cells, "role"),
            "partner": find_col(cells, "partner"), "repeat": find_col(cells, "repeat")}
    missing = [k for k in ("pair", "id", "role") if cols[k] is None]
    if missing:
        raise NPSError("draft_layout", f"The Pairing sheet is missing these columns: {missing}. Restore the header row.")
    data = []
    for rownum, r in enumerate(rows[i + 1:], start=i + 2):
        if clean(cell(r, cols["id"])):
            data.append({"row": rownum, "pair": clean(cell(r, cols["pair"])), "id": clean(cell(r, cols["id"])),
                         "key": norm_id(cell(r, cols["id"])), "role": clean(cell(r, cols["role"]))})
    return i + 1, cols, data


def validate_draft(paths: Paths, roster=None, history=None, settings=None) -> dict:
    run_dir = paths.require_run()
    info = read_info(run_dir)
    settings = settings or load_settings(paths, info["Simulation"])
    roster = roster or load_roster(paths)
    history = history or load_history(paths)
    by_key = {s["key"]: s for s in roster["pairable"]}
    name = lambda k: by_key[k]["name"] if k in by_key else roster["by_key"].get(k, {}).get("name", k)
    absent = {norm_id(x) for x in info.get("Absent IDs", "").split(",") if x.strip()}
    odd, odd_student = info.get("Odd Arrangement", ""), norm_id(info.get("Odd Student ID", ""))
    apart = {frozenset(norm_id(x) for x in p.split(",")) for p in info.get("Keep Apart", "").split(";") if p.strip()}
    roles = settings["roles"]
    role_lookup = {r.lower(): r for r in roles}

    path = run_dir / "Pairing Draft.xlsx"
    try:
        wb = load_workbook(path)
    except PermissionError:
        raise NPSError("file_locked", "Pairing Draft.xlsx is open in Excel or locked by OneDrive. Close it and try again.")
    ws = wb[PAIRING_SHEET] if PAIRING_SHEET in wb.sheetnames else wb.worksheets[0]
    header_row, cols, rows = read_draft_rows(ws)
    errors, warnings = [], []

    counts = Counter(r["key"] for r in rows)
    for k, n in counts.items():
        if n > 1:
            errors.append(f"{name(k)} appears {n} times.")
    attending = {s["key"] for s in roster["pairable"]} - absent
    for k in sorted(attending - set(counts), key=name):
        errors.append(f"{name(k)} is attending but is not in the draft.")
    for k in counts:
        if k in absent:
            errors.append(f"{name(k)} was marked absent but is in the draft.")
        elif k not in by_key:
            errors.append(f"Student ID {k} is not an enrolled student on the roster.")

    groups = defaultdict(list)
    for r in rows:
        role = role_lookup.get(r["role"].lower())
        if r["role"].lower() == OBSERVER.lower():
            if not (odd == "observe" and r["key"] == odd_student):
                errors.append(f"{name(r['key'])} is marked Observer, but no observer was chosen for this round.")
            continue
        if role is None:
            errors.append(f"{name(r['key'])} has role '{r['role']}', which is not {roles[0]} or {roles[1]}.")
        if not r["pair"]:
            errors.append(f"{name(r['key'])} has no pair number.")
            continue
        groups[r["pair"]].append(r)

    partner_text, repeat_text = {}, {}
    for pair, members in groups.items():
        keys = [m["key"] for m in members]
        member_roles = Counter((role_lookup.get(m["role"].lower()) or m["role"]) for m in members)
        size = len(members)
        if size == 1:
            if not (odd == "instructor" and keys[0] == odd_student):
                errors.append(f"Pair {pair} has only one student ({name(keys[0])}).")
        elif size == 3:
            if odd != "trio":
                errors.append(f"Pair {pair} has three students, but no group of three was chosen for this round.")
            elif set(member_roles) != set(roles):
                errors.append(f"Group {pair} needs at least one {roles[0]} and one {roles[1]}.")
        elif size > 3:
            errors.append(f"Pair {pair} has {size} students.")
        elif set(member_roles) != set(roles):
            errors.append(f"Pair {pair} ({' and '.join(name(k) for k in keys)}) does not have one {roles[0]} and one {roles[1]}.")
        for m in members:
            others = [name(k) for k in keys if k != m["key"]]
            partner_text[m["row"]] = ", ".join(others) if others else f"{INSTRUCTOR} (you)"
        for a_, b_ in combinations(keys, 2):
            shared = history["shared"].get(frozenset((a_, b_)))
            if shared:
                msg = f"{name(a_)} and {name(b_)} already negotiated together ({', '.join(shared)})."
                warnings.append("Repeat partner: " + msg)
                for m in members:
                    if m["key"] in (a_, b_):
                        repeat_text[m["row"]] = "Yes: " + msg
            if frozenset((a_, b_)) in apart:
                warnings.append(f"Kept apart: you asked to keep {name(a_)} and {name(b_)} apart, but they are paired.")

    if settings["balancing"] and history["tracks_roles"]:
        for r in rows:
            role = role_lookup.get(r["role"].lower())
            entry = history["by_key"].get(r["key"])
            if not role or not entry:
                continue
            other = roles[1] if role == roles[0] else roles[0]
            had, had_other = entry["roles"].get(role.lower(), 0), entry["roles"].get(other.lower(), 0)
            if had > had_other:
                times = lambda n: f"{n} time" + ("" if n == 1 else "s")
                warnings.append(f"Role balance: {name(r['key'])} is {role} again "
                                f"(has been {role} {times(had)} and {other} {times(had_other)}).")

    if cols["partner"] is not None:
        for r in rows:
            ws.cell(r["row"], cols["partner"] + 1, partner_text.get(r["row"], ""))
    if cols["repeat"] is not None:
        for r in rows:
            ws.cell(r["row"], cols["repeat"] + 1, repeat_text.get(r["row"], ""))

    checks = wb[CHECKS_SHEET] if CHECKS_SHEET in wb.sheetnames else wb.create_sheet(CHECKS_SHEET)
    checks.delete_rows(1, checks.max_row)
    summary = f"{len(errors)} errors, {len(warnings)} warnings (checked {datetime.now():%Y-%m-%d %H:%M})"
    checks.append(["Level", "Finding"])
    style_header_row(checks)
    checks.append(["Summary", summary])
    for e in errors:
        checks.append(["Error (must fix before approving)", e])
    for w in warnings:
        checks.append(["Warning (can approve if you accept it)", w])
    if not errors and not warnings:
        checks.append(["OK", "Every attending student appears once, every pair has both roles, and no one repeats a partner."])
    checks.column_dimensions["A"].width, checks.column_dimensions["B"].width = 38, 110
    save_workbook(wb, path)
    return {"errors": errors, "warnings": warnings, "summary": summary}


def cmd_validate(a):
    paths = Paths(a.root, a.section, a.run)
    checks = validate_draft(paths)
    finish({"ok": True, "run": a.run, **checks}, paths, "T4 Validate draft", checks["summary"])


# ---------------------------------------------------------------- T5 record

def cmd_record(a):
    paths = Paths(a.root, a.section, a.run)
    run_dir = paths.require_run()
    info = read_info(run_dir)
    if info.get("Status", "").startswith("Recorded"):
        raise NPSError("already_recorded", f"This pairing was already recorded ({info['Status']}).")
    settings = load_settings(paths, info["Simulation"])
    roster, history = load_roster(paths), load_history(paths)
    checks = validate_draft(paths, roster, history, settings)
    if checks["errors"]:
        raise NPSError("draft_has_errors", "The draft has errors, so it cannot be recorded.", checks)
    if checks["warnings"] and not a.accept_warnings:
        raise NPSError("warnings_not_accepted", "The draft has warnings. Confirm them with the instructor, "
                       "then run record again with --accept-warnings.", checks)

    when = date.fromisoformat(info["Negotiation Date"])
    label = f"{when.month}/{when.day} {info['Simulation']}"
    pair_label, role_label = f"{label} Pair ID", f"{label} Role"
    if any(low(h) == pair_label.lower() for h in history["cells"]):
        raise NPSError("already_recorded", f"Pairing History already has a '{pair_label}' column.")

    wb_draft = load_workbook(run_dir / "Pairing Draft.xlsx")
    _, _, rows = read_draft_rows(wb_draft[PAIRING_SHEET])
    groups = defaultdict(list)
    for r in rows:
        if r["role"].lower() != OBSERVER.lower() and r["pair"]:
            groups[r["pair"]].append(r)
    real_groups = [g for _, g in sorted(groups.items(), key=lambda x: pair_sort_key(x[0])) if len(g) > 1]
    ids = new_pair_ids(history["format"], len(real_groups))
    assigned = {m["key"]: (pid, m["role"]) for pid, g in zip(ids, real_groups) for m in g}
    solo = {g[0]["key"]: ("", g[0]["role"]) for g in groups.values() if len(g) == 1}
    track_roles = history["tracks_roles"] or (not history["rounds"] and settings["balancing"])

    backup_name = paths.backup(paths.history, f"before {info['Simulation']}")
    try:
        wb = load_workbook(paths.history)
    except PermissionError:
        raise NPSError("file_locked", "Pairing History.xlsx is open in Excel or locked by OneDrive. Close it and try again.")
    ws = wb.worksheets[0]
    hr = history["header_row"]
    last_col = max((c.column for c in ws[hr] if c.value is not None), default=0)
    pair_col, role_col = last_col + 1, last_col + 2
    ws.cell(hr, pair_col, pair_label)
    if track_roles:
        ws.cell(hr, role_col, role_label)
    template = ws.cell(hr, last_col) if last_col else None
    for col in ([pair_col, role_col] if track_roles else [pair_col]):
        target = ws.cell(hr, col)
        if template is not None and template.has_style:
            target.font, target.fill = copy(template.font), copy(template.fill)
            target.border, target.alignment = copy(template.border), copy(template.alignment)
        else:
            target.font = HEADER_FONT
        ws.column_dimensions[target.column_letter].width = max(14, len(str(target.value)) + 2)

    c_id = find_col(history["cells"], "student id")
    c_name = find_col(history["cells"], "name")
    rows_by_key = {e["key"]: e["row"] for e in history["entries"]}
    by_key = {s["key"]: s for s in roster["pairable"]}
    added_rows = []
    for key, (pid, role) in {**assigned, **solo}.items():
        row = rows_by_key.get(key)
        if row is None:
            row = max([hr] + list(rows_by_key.values())) + 1
            rows_by_key[key] = row
            if c_name is not None:
                ws.cell(row, c_name + 1, by_key[key]["name"])
            id_cell = ws.cell(row, c_id + 1, by_key[key]["id"])
            id_cell.number_format = "@"
            added_rows.append(by_key[key]["name"])
        if pid:
            ws.cell(row, pair_col, pid)
        if track_roles:
            ws.cell(row, role_col, role)
    save_workbook(wb, paths.history)

    check = load_history(paths)
    col_index = next(rd["pair_col"] for rd in check["rounds"] if rd["label"] == pair_label)
    for key, (pid, _) in assigned.items():
        if check["by_key"].get(key, {}).get("ids", {}).get(col_index) != pid:
            shutil.copy2(paths.backups / backup_name, paths.history)
            raise NPSError("record_mismatch", "Pairing History did not save correctly, so the backup was restored.")

    info["Status"] = f"Recorded {datetime.now():%Y-%m-%d %H:%M}"
    info["Pair IDs"] = f"{ids[0]} to {ids[-1]}" if ids else ""
    write_info(wb_draft[INFO_SHEET], info)
    save_workbook(wb_draft, run_dir / "Pairing Draft.xlsx")
    finish({"ok": True, "run": a.run, "pair_ids": info["Pair IDs"], "columns_added": [pair_label] +
            ([role_label] if track_roles else []), "students_recorded": len(assigned) + len(solo),
            "rows_added_to_history": added_rows, "backup": backup_name,
            "accepted_warnings": checks["warnings"]},
           paths, "T5 Record pairing",
           f"Approved by instructor. Recorded {info['Pair IDs']} in columns '{pair_label}'"
           f"{' and ' + repr(role_label) if track_roles else ''}. Backup: {backup_name}. "
           f"Accepted warnings: {len(checks['warnings'])}.")


# ---------------------------------------------------------------- T6 mail merge

def cmd_mailmerge(a):
    paths = Paths(a.root, a.section, a.run)
    run_dir = paths.require_run()
    info = read_info(run_dir)
    if not info.get("Status", "").startswith("Recorded"):
        raise NPSError("not_recorded", "Record the approved pairing in Pairing History before building the mail merge.")
    settings = load_settings(paths, info["Simulation"])
    roster = load_roster(paths)
    by_key = roster["by_key"]
    missing_info = [r for r, text in settings["role_info"].items() if not text]
    if missing_info:
        raise NPSError("missing_role_info", f"Simulation Settings has no role information for: {missing_info}.")
    _, _, rows = read_draft_rows(load_workbook(run_dir / "Pairing Draft.xlsx")[PAIRING_SHEET])
    role_lookup = {r.lower(): r for r in settings["roles"]}
    wb = Workbook()
    ws = wb.active
    ws.title = "Mail Merge"
    ws.append(MAIL_MERGE_HEADERS)
    expected, missing_email, observers = {}, [], []
    for r in rows:
        if r["role"].lower() == OBSERVER.lower():
            observers.append(by_key[r["key"]]["name"])
            continue
        s, role = by_key[r["key"]], role_lookup[r["role"].lower()]
        status = "Ready" if s["email"] else "Missing email"
        if not s["email"]:
            missing_email.append(s["name"])
        ws.append([s["first"], s["last"], s["email"], info["Simulation"], role, settings["role_info"][role],
                   settings["survey_link"], status])
        expected[s["email"] or s["name"]] = role
    for col, width in zip("ABCDEFGH", (14, 16, 32, 14, 12, 60, 50, 14)):
        ws.column_dimensions[col].width = width
    style_header_row(ws)
    path = run_dir / "Role Mail Merge.xlsx"
    save_workbook(wb, path)
    back = {clean(r[2]) or f"{clean(r[0])} {clean(r[1])}": clean(r[4])
            for r in load_workbook(path).active.iter_rows(min_row=2, values_only=True)}
    if back != expected:
        raise NPSError("mailmerge_mismatch", "Role Mail Merge.xlsx did not match the approved pairing after saving.")
    finish({"ok": True, "run": a.run, "workbook": str(path), "rows": len(expected), "missing_email": missing_email,
            "observers_not_emailed": observers, "columns": MAIL_MERGE_HEADERS},
           paths, "T6 Build mail merge",
           f"{len(expected)} rows. Missing email: {missing_email or 'none'}. Observers: {observers or 'none'}.")


# ---------------------------------------------------------------- T7/T8 survey

def read_survey(path: Path, settings: dict):
    if not path.exists():
        raise NPSError("file_missing", f"Could not find {path.name} in {path.parent.name}. Save the Qualtrics export "
                                       "there as 'Survey Export.csv'.")
    try:
        with path.open(newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    except PermissionError:
        raise NPSError("file_locked", f"{path.name} is open in Excel. Close it and try again.")
    if not rows:
        raise NPSError("survey_empty", "Survey Export.csv is empty.")
    codes = [c.strip() for c in rows[0]]
    skip = 0
    if len(rows) > 2 and rows[2] and rows[2][0].startswith('{"ImportId'):
        skip = 2
    elif len(rows) > 1 and rows[1] and rows[1][0].startswith('{"ImportId'):
        skip = 1
    texts = [c.strip() for c in rows[1]] if skip == 2 else codes

    def col(name):
        for header in (codes, texts):
            for i, h in enumerate(header):
                if h.lower() == name.lower():
                    return i
        return None

    c_id = col(settings["survey_id_column"])
    if c_id is None:
        raise NPSError("survey_column_missing", f"The export has no '{settings['survey_id_column']}' column for student IDs.",
                       {"columns": codes})
    mapped = []
    for src, dst in settings["mapping"]:
        i = col(src)
        if i is None:
            raise NPSError("survey_column_missing", f"The export has no '{src}' column (mapped to '{dst}').",
                           {"columns": codes})
        mapped.append((i, dst))
    c_status, c_resp = col("Status"), col("ResponseId")
    c_time = col("RecordedDate") if col("RecordedDate") is not None else col("EndDate")
    responses, removed = [], []
    for n, r in enumerate(rows[1 + skip:], start=1):
        if not any(x.strip() for x in r):
            continue
        status = cell(r, c_status) or ""
        resp_id = (cell(r, c_resp) or f"row{n}").strip()
        if any(word in status.lower() for word in ("preview", "test", "spam")):
            removed.append({"response": resp_id, "reason": f"Status '{status}'"})
            continue
        responses.append({"response": resp_id, "id_entered": (cell(r, c_id) or "").strip(),
                          "key": norm_id(cell(r, c_id)), "time": (cell(r, c_time) or "").strip(),
                          "values": {dst: parse_number(cell(r, i)) for i, dst in mapped}})
    return responses, removed


def survey_match(paths: Paths, a):
    run_dir = paths.require_run()
    info = read_info(run_dir)
    if not info.get("Status", "").startswith("Recorded"):
        raise NPSError("not_recorded", "This simulation's pairing has not been approved and recorded yet.")
    settings = load_settings(paths, info["Simulation"])
    roster = load_roster(paths)
    _, _, rows = read_draft_rows(load_workbook(run_dir / "Pairing Draft.xlsx")[PAIRING_SHEET])
    paired = [r for r in rows if r["role"].lower() != OBSERVER.lower()]
    paired_keys = {r["key"] for r in paired}
    responses, removed = read_survey(run_dir / "Survey Export.csv", settings)
    by_resp = {r["response"]: r for r in responses}
    name = lambda k: roster["by_key"].get(k, {}).get("name", k)

    decisions_used = []
    for resp in a.leave_out:
        if resp not in by_resp:
            raise NPSError("unknown_response", f"There is no response {resp} in the export.")
        by_resp[resp]["left_out"] = True
        decisions_used.append(f"Left out {resp}")
    for spec in a.assign:
        resp, _, sid = spec.partition("=")
        if resp not in by_resp or norm_id(sid) not in paired_keys:
            raise NPSError("bad_assign", f"--assign {spec}: response or student ID not found in this pairing.")
        by_resp[resp]["key"], by_resp[resp]["assigned"] = norm_id(sid), True
        decisions_used.append(f"Assigned {resp} to {name(norm_id(sid))}")

    per_student = defaultdict(list)
    unknown = []
    for r in responses:
        if r.get("left_out"):
            continue
        (per_student[r["key"]] if r["key"] in paired_keys else unknown).append(r)

    chosen, exceptions, notes = {}, [], {}
    for k in paired_keys:
        subs = sorted(per_student.get(k, []), key=lambda r: r["time"])
        if not subs:
            continue
        if len(subs) == 1:
            chosen[k] = subs[0]
            if subs[0].get("assigned"):
                notes[k] = f"Response matched by instructor (ID entered: {subs[0]['id_entered']})"
            continue
        pick = [s for s in subs if s["response"] in a.use]
        if len(pick) == 1:
            chosen[k] = pick[0]
            notes[k] = f"Used response {pick[0]['response']} of {len(subs)} submissions (instructor's choice)"
        elif a.latest_duplicates:
            chosen[k] = subs[-1]
            notes[k] = f"Used latest of {len(subs)} submissions"
        else:
            exceptions.append({"type": "duplicate", "student": name(k), "student_id": roster["by_key"][k]["id"],
                               "submissions": [{"response": s["response"], "time": s["time"], "values": s["values"]}
                                               for s in subs]})
    no_response_keys = [k for k in paired_keys if k not in chosen and not per_student.get(k)]
    for r in unknown:
        hint = None
        if r["key"] in roster["by_key"]:
            hint = f"This ID belongs to {name(r['key'])}, who was not paired this round."
        candidates = sorted(paired_keys, key=lambda k: -difflib.SequenceMatcher(None, r["key"], k).ratio())
        candidates = sorted(candidates[:3], key=lambda k: k not in no_response_keys)
        exceptions.append({"type": "unknown_id", "response": r["response"], "id_entered": r["id_entered"],
                           "time": r["time"], "values": r["values"], **({"hint": hint} if hint else {}),
                           "suggestions": [{"student": name(k), "student_id": roster["by_key"][k]["id"],
                                            "has_no_response": k in no_response_keys} for k in candidates]})
    for k in no_response_keys:
        notes[k] = "No survey response"
    if a.latest_duplicates and any("latest" in n for n in notes.values()):
        decisions_used.append("Used the latest submission for every duplicate")
    for resp in a.use:
        decisions_used.append(f"Used response {resp} for its student's duplicate submissions")
    return {"info": info, "settings": settings, "rows": paired, "roster": roster, "chosen": chosen,
            "notes": notes, "exceptions": exceptions, "removed": removed, "decisions": decisions_used,
            "no_response": [name(k) for k in no_response_keys], "response_count": len(responses)}


def cmd_survey(a):
    paths = Paths(a.root, a.section, a.run)
    m = survey_match(paths, a)
    finish({"ok": True, "run": a.run, "responses": m["response_count"], "matched": len(m["chosen"]),
            "no_response": m["no_response"], "ignored_rows": m["removed"], "exceptions": m["exceptions"],
            "ready_for_outcomes": not m["exceptions"], "decisions_applied": m["decisions"]},
           paths, "T8 Match survey",
           f"{m['response_count']} responses, {len(m['chosen'])} matched, {len(m['no_response'])} no response, "
           f"{len(m['exceptions'])} exceptions, {len(m['removed'])} ignored rows.")


# ---------------------------------------------------------------- T9 outcomes

def cmd_outcomes(a):
    started = time.time()
    paths = Paths(a.root, a.section, a.run)
    m = survey_match(paths, a)
    if m["exceptions"]:
        raise NPSError("unresolved_exceptions", "Some survey responses need the instructor's decision first.",
                       {"exceptions": m["exceptions"]})
    settings, info = m["settings"], m["info"]
    if not settings["template"]:
        raise NPSError("template_missing", f"Simulation Settings has no outcomes template for {info['Simulation']}.")
    template = paths.templates / settings["template"]
    if not template.exists():
        raise NPSError("template_missing", f"Could not find the template {template.name} in the Templates folder.")
    try:
        wb = load_workbook(template)
    except PermissionError:
        raise NPSError("file_locked", f"{template.name} is open in Excel. Close it and try again.")
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi, cells = find_header(rows, ["role"])
    if hi is None:
        raise NPSError("template_layout", f"{template.name} has no header row with a 'Role' column.")
    cols = {"pair": find_col(cells, "pair"), "name": find_col(cells, "name"), "id": find_col(cells, "student id"),
            "role": find_col(cells, "role"), "notes": find_col(cells, "note")}
    value_cols = {}
    for _, dst in settings["mapping"]:
        j = next((i for i, h in enumerate(cells) if h == dst.lower()), None)
        if j is None:
            raise NPSError("template_column_missing", f"{template.name} has no '{dst}' column.")
        value_cols[dst] = j

    roles = settings["roles"]
    order = {r.lower(): i for i, r in enumerate(roles)}
    groups = defaultdict(list)
    for r in m["rows"]:
        groups[r["pair"]].append(r)
    ordered = [r for _, g in sorted(groups.items(), key=lambda x: pair_sort_key(x[0]))
               for r in sorted(g, key=lambda r: order.get(r["role"].lower(), 9))]
    by_key = m["roster"]["by_key"]
    expected = {}
    for offset, r in enumerate(ordered, start=1):
        row = hi + 1 + offset
        s = by_key[r["key"]]
        for key, value in (("pair", parse_number(r["pair"])), ("name", s["name"]), ("id", s["id"]), ("role", r["role"])):
            if cols[key] is not None:
                c = ws.cell(row, cols[key] + 1, value)
                if key == "id":
                    c.number_format = "@"
        response = m["chosen"].get(r["key"])
        for dst, j in value_cols.items():
            value = response["values"][dst] if response else None
            ws.cell(row, j + 1, value)
            expected[(row, j + 1)] = value
        if cols["notes"] is not None and r["key"] in m["notes"]:
            ws.cell(row, cols["notes"] + 1, m["notes"][r["key"]])

    out = paths.run_dir / "Outcomes.xlsx"
    previous = paths.backup(out, "replaced") if out.exists() else None
    tmp = paths.run_dir / "~Outcomes building.xlsx"
    save_workbook(wb, tmp)
    check = load_workbook(tmp).worksheets[0]
    mismatches = [f"{check.cell(r, c).coordinate}" for (r, c), v in expected.items() if check.cell(r, c).value != v]
    if mismatches:
        tmp.unlink()
        raise NPSError("outcomes_mismatch", "Some survey values did not save correctly; the workbook was not created.",
                       {"cells": mismatches})
    if previous:
        out.unlink()
    tmp.replace(out)
    seconds = round(time.time() - started, 1)
    finish({"ok": True, "run": a.run, "workbook": str(out), "students": len(ordered), "matched": len(m["chosen"]),
            "no_response": m["no_response"], "decisions_applied": m["decisions"], "ignored_rows": m["removed"],
            "previous_outcomes_backup": previous, "seconds": seconds},
           paths, "T9 Build outcomes",
           f"Outcomes.xlsx: {len(ordered)} students, {len(m['chosen'])} with survey data, "
           f"{len(m['no_response'])} no response. Decisions: {m['decisions'] or 'none'}. "
           f"Previous file backup: {previous or 'none'}.")


# ---------------------------------------------------------------- log

def cmd_log(a):
    paths = Paths(a.root, a.section, a.run)
    finish({"ok": True}, paths, a.step, a.details)


# ---------------------------------------------------------------- CLI

def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="command", required=True)

    def add(name, func, run=False, simulation=False, survey=False):
        sp = sub.add_parser(name)
        sp.add_argument("--root", required=True, help="Course folder containing Simulation Settings.xlsx")
        sp.add_argument("--section", required=True, help="Section folder name, e.g. 'Fall 2026 Section 01'")
        if run:
            sp.add_argument("--run", required=True, help="Simulation folder name, e.g. 'Used Car 2026-10-14'")
        if simulation:
            sp.add_argument("--simulation", required=True)
        if survey:
            sp.add_argument("--latest-duplicates", action="store_true", help="Use the latest of duplicate submissions")
            sp.add_argument("--use", action="append", default=[], help="Response ID to use for a duplicate")
            sp.add_argument("--assign", action="append", default=[], help="RESPONSE_ID=STUDENT_ID for an unknown ID")
            sp.add_argument("--leave-out", action="append", default=[], help="Response ID to leave out")
        sp.set_defaults(func=func)
        return sp

    add("snapshot", cmd_snapshot, simulation=True)
    sp = add("match-absences", cmd_match_absences)
    sp.add_argument("--absent", action="append", default=[], help="One absence entry as the instructor wrote it")
    sp = add("pair", cmd_pair, simulation=True)
    sp.add_argument("--date", help="Negotiation date YYYY-MM-DD (default today)")
    sp.add_argument("--absent-id", action="append", default=[], help="Student ID of an absent student")
    sp.add_argument("--odd", choices=["trio", "instructor", "observe"])
    sp.add_argument("--odd-student-id")
    sp.add_argument("--trio-role", help="Role given to two members of the group of three")
    sp.add_argument("--keep-apart", action="append", default=[], help="ID1,ID2 of students to keep apart")
    sp.add_argument("--seed", type=int)
    add("validate", cmd_validate, run=True)
    sp = add("record", cmd_record, run=True)
    sp.add_argument("--accept-warnings", action="store_true")
    add("mailmerge", cmd_mailmerge, run=True)
    add("survey", cmd_survey, run=True, survey=True)
    add("outcomes", cmd_outcomes, run=True, survey=True)
    sp = add("log", cmd_log)
    sp.add_argument("--run")
    sp.add_argument("--step", required=True)
    sp.add_argument("--details", required=True)

    a = p.parse_args()
    try:
        a.func(a)
    except NPSError as exc:
        try:  # record the failure in the Run Log too, when the section folder exists
            write_log(Paths(a.root, a.section, getattr(a, "run", None)), f"{a.command} failed", exc.message)
        except Exception:
            pass
        print(json.dumps({"ok": False, "error": exc.code, "message": exc.message, **exc.details},
                         indent=2, ensure_ascii=False))
        sys.exit(2)


if __name__ == "__main__":
    main()

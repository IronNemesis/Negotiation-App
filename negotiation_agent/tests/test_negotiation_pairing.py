"""Tests for the negotiation-pairing skill, run against the fictional example files.

    python negotiation_agent/tests/test_negotiation_pairing.py

Works with pytest too. Each test copies the examples into a temporary folder, so the repo is never changed.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = ROOT / "examples"
SCRIPT = ROOT / "skills" / "negotiation-pairing" / "scripts" / "pairing.py"
sys.path.insert(0, str(SCRIPT.parent))
import nps_shared as ns  # noqa: E402

BEFORE = "BUS 4489 F26 Roster - before Used Car.xlsx"
DRAFTED = "BUS 4489 F26 Roster - Used Car drafted.xlsx"
AFTER = "BUS 4489 F26 Roster - after Used Car.xlsx"
EXPECTED = json.loads((EXAMPLES / "expected" / "used-car.json").read_text(encoding="utf-8"))
ABSENT = ["Adebayo", "Castellano, Rosa", "tzimmerm@example.edu"]  # three ways she might name them


def workspace() -> Path:
    tmp = Path(tempfile.mkdtemp())
    for f in (BEFORE, DRAFTED, AFTER):
        shutil.copy(EXAMPLES / f, tmp / f)
    return tmp


def run(folder: Path, *args) -> dict:
    out = subprocess.run([sys.executable, str(SCRIPT), *args], cwd=folder, capture_output=True, text=True)
    assert out.stdout, out.stderr
    return json.loads(out.stdout)


def as_tables(groups: dict) -> dict:
    return {m: {r: [p.name for p in ps if p.role == r] for r in ("BUYER", "SELLER")} for m, ps in groups.items()}


def expected_tables(key: str) -> dict:
    return {m["match"]: {"BUYER": m["BUYER"], "SELLER": m["SELLER"]} for m in EXPECTED[key]}


def uc_tab(path: Path) -> ns.SimTab:
    return ns.read_sim_tab(load_workbook(path), "UC", ["BUYER", "SELLER"])


def all_values(path: Path, skip_tab: str | None = None) -> dict:
    wb = load_workbook(path)
    return {(ws.title, c.coordinate): c.value for ws in wb.worksheets if ws.title != skip_tab
            for row in ws.iter_rows() for c in row if c.value is not None}


def test_draft_preview_uses_her_alphabetical_method():
    result = run(workspace(), "draft", "--roster", BEFORE)
    assert result["ok"] and result["matches"] == 19 and not result["errors"]
    preview = {m["match"]: {"BUYER": m["BUYER"], "SELLER": m["SELLER"]} for m in result["draft"]}
    assert preview == expected_tables("draft")


def test_write_draft_fills_only_the_draft_table():
    folder = workspace()
    before = all_values(folder / BEFORE, skip_tab="UC")
    result = run(folder, "write-draft", "--roster", BEFORE)
    assert result["ok"], result
    assert Path(result["backup"]).exists()
    assert all_values(folder / BEFORE, skip_tab="UC") == before, "a tab other than UC changed"
    tab = uc_tab(folder / BEFORE)
    assert as_tables(tab.groups(tab.draft_members)) == expected_tables("draft")
    assert not tab.final_members
    # The draft table now matches the example "drafted" workbook cell for cell.
    ours, theirs = load_workbook(folder / BEFORE)["UC"], load_workbook(folder / DRAFTED)["UC"]
    for r in range(tab.draft_header, tab.draft_header + 23):
        assert [ours.cell(r, c).value for c in (1, 2, 3)] == [theirs.cell(r, c).value for c in (1, 2, 3)], r
    again = run(folder, "write-draft", "--roster", BEFORE)
    assert not again["ok"] and again["error"] == "draft_exists"
    assert run(folder, "write-draft", "--roster", BEFORE, "--replace")["ok"]


def test_known_absence_and_swap():
    folder = workspace()
    result = run(folder, "draft", "--roster", BEFORE, "--absent", "Owen Bennett",
                 "--swap", "Alvarez, Maya;Nakamura, Ren")
    assert result["ok"] and result["known_absent"] == ["Bennett, Owen"]
    names = [n for m in result["draft"] for r in ("BUYER", "SELLER") for n in m[r]]
    assert "Bennett, Owen" not in names and len(names) == 38 and result["matches"] == 19
    # Without Owen Bennett the halves shift: Maya Alvarez (a Buyer) and Ren Nakamura (now the first Seller)
    # trade places, keeping each other's role and match.
    role_of = {n: r for m in result["draft"] for r in ("BUYER", "SELLER") for n in m[r]}
    assert role_of["Alvarez, Maya"] == "SELLER" and role_of["Nakamura, Ren"] == "BUYER"
    unclear = run(folder, "draft", "--roster", BEFORE, "--absent", "Park")
    assert not unclear["ok"] and unclear["error"] == "absence_unclear"
    assert set(unclear["unclear"][0]["suggestions"]) == {"Park, Daniel", "Park, Danielle"}


def test_check_reports_her_excel_edits():
    folder = workspace()
    wb = load_workbook(folder / DRAFTED)
    ws = wb["UC"]
    tab = ns.read_sim_tab(wb, "UC", ["BUYER", "SELLER"])
    row = tab.draft_header + 2                      # match 2
    ws.cell(row, 3).value = ws.cell(row - 1, 3).value  # her typo: match 1's Seller copied into match 2
    wb.save(folder / DRAFTED)
    result = run(folder, "check", "--roster", DRAFTED)
    assert result["ok"]
    assert any("appears twice" in e for e in result["errors"])
    assert any("Nakamura, Ren" in w and "not in any match" in w for w in result["warnings"])


def test_bcc_lists_by_role():
    folder = workspace()
    result = run(folder, "bcc", "--roster", DRAFTED)
    assert result["ok"]
    assert result["lists"]["BUYER"]["count"] == 19 and result["lists"]["SELLER"]["count"] == 20
    assert "tzimmerm@example.edu" in result["lists"]["SELLER"]["bcc"]
    assert Path(result["saved_to"]).exists()


def test_absences_follow_her_rule_and_match_the_after_example():
    folder = workspace()
    args = ["--roster", DRAFTED, *sum((["--absent", a] for a in ABSENT), [])]
    proposal = run(folder, "absences", *args)
    assert proposal["ok"] and not proposal["errors"], proposal
    final = {m["match"]: {"BUYER": m["BUYER"], "SELLER": m["SELLER"]} for m in proposal["final"]}
    assert final == expected_tables("final")
    assert any("Mendes, Rafael joins match 2" in m for m in proposal["moves"])
    assert any("O'Neill, Siobhan joins match 4" in m for m in proposal["moves"])
    before = all_values(folder / DRAFTED, skip_tab="UC")
    written = run(folder, "write-final", *args)
    assert written["ok"], written
    assert all_values(folder / DRAFTED, skip_tab="UC") == before
    tab = uc_tab(folder / DRAFTED)
    assert as_tables(tab.groups(tab.final_members)) == expected_tables("final")
    assert {m.name: m.pair_id for m in tab.absent} == {n: "UC_0" for n in EXPECTED["absent_in_class"]}
    # The whole final table (names, emails, pair IDs) matches the example "after" workbook.
    ours, theirs = load_workbook(folder / DRAFTED)["UC"], load_workbook(folder / AFTER)["UC"]
    for r in range(1, tab.draft_header):
        assert [ours.cell(r, c).value for c in range(1, 11)] == [theirs.cell(r, c).value for c in range(1, 11)], r
    again = run(folder, "write-final", *args)
    assert not again["ok"] and again["error"] == "final_area_not_empty"


def test_stranded_student_who_doesnt_fit_needs_her_choice():
    folder = workspace()
    # Grace Mahoney is the only Buyer in match 19 (with two Sellers): both Sellers are stranded, and
    # match 18 can take only one of them.
    proposal = run(folder, "absences", "--roster", DRAFTED, "--absent", "Mahoney")
    assert proposal["ok"] and proposal["needs_placement"], proposal
    stuck = proposal["needs_placement"][0]["student"]
    placed = run(folder, "absences", "--roster", DRAFTED, "--absent", "Mahoney", "--place", f"{stuck}=17")
    assert placed["ok"] and not placed["errors"] and not placed["needs_placement"], placed


def drop_from_class_list(path: Path, name: str) -> None:
    wb = load_workbook(path)
    ws = wb["Att - Fall"]
    row = next(c.row for c in ws["B"] if c.value == name)
    ws.delete_rows(row)
    wb.save(path)


def test_student_who_dropped_after_the_draft_is_treated_as_absent():
    folder = workspace()
    drop_from_class_list(folder / DRAFTED, "Bennett, Owen")  # the only Buyer in match 3
    check = run(folder, "check", "--roster", DRAFTED)
    assert check["ok"] and not check["errors"], check
    assert any("Bennett, Owen" in w and "listed as absent" in w for w in check["warnings"])
    # Not named as absent: still listed as absent, and the partner moves under her rule.
    for named in ([], ["--absent", "Bennett"]):
        proposal = run(folder, "absences", "--roster", DRAFTED, *named)
        assert proposal["ok"] and not proposal["errors"], proposal
        assert proposal["absent"] == ["Bennett, Owen"] and proposal["dropped"] == ["Bennett, Owen"]
        assert any("Nguyen, Katherine joins match 2" in m and "no longer on the class list" in m
                   for m in proposal["moves"])
    written = run(folder, "write-final", "--roster", DRAFTED)
    assert written["ok"], written
    tab = uc_tab(folder / DRAFTED)
    assert [(m.name, m.role, m.pair_id) for m in tab.absent] == [("Bennett, Owen", "BUYER", "UC_0")]
    assert 3 not in tab.groups(tab.final_members)


def test_a_likely_name_is_still_asked_about_not_treated_as_dropped():
    folder = workspace()
    wb = load_workbook(folder / DRAFTED)
    ws = wb["Att - Fall"]
    next(c for c in ws["B"] if c.value == "Bennett, Owen").value = "Bennett, Oliver"  # same last name, new first
    wb.save(folder / DRAFTED)
    proposal = run(folder, "absences", "--roster", DRAFTED)
    assert proposal["ok"] and not proposal["dropped"]
    assert any("may be Bennett, Oliver" in e for e in proposal["errors"]), proposal["errors"]


def write_sl_draft(path: Path, rows: list[tuple]) -> None:
    """Fill the SL draft table by hand, as she might: (match, supervisor, subordinate) per row."""
    wb = load_workbook(path)
    tab = ns.read_sim_tab(wb, "SL", ["SUPERVISOR", "SUBORDINATE"])
    ws = wb["SL"]
    for i, (match, sup, sub) in enumerate(rows):
        r = tab.draft_header + 1 + i
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value = match, sup, sub
    wb.save(path)


def test_check_warns_when_someone_is_the_one_in_a_1v2_again():
    folder = workspace()
    # In Used Car, Maya Alvarez was the only Buyer in match 2, facing two Sellers.
    write_sl_draft(folder / AFTER, [(1, "Alvarez, Maya", "Shah, Dev"), (None, None, "Silva, Kai")])
    result = run(folder, "check", "--roster", AFTER, "--simulation", "SL")
    assert result["ok"]
    assert any(w.startswith("Alvarez, Maya faces two partners alone in match 1") and "UC_2" in w
               for w in result["warnings"]), result["warnings"]


def test_absence_update_asks_before_making_someone_the_one_again():
    folder = workspace()
    # Tyler Brooks was the only Buyer in Used Car match 4. If Ava Quinn is absent, Noor Reyes would join match 1
    # and leave Tyler facing two partners again, so she is asked instead.
    write_sl_draft(folder / AFTER, [(1, "Brooks, Tyler", "Shah, Dev"), (2, "Quinn, Ava", "Reyes, Noor")])
    proposal = run(folder, "absences", "--roster", AFTER, "--simulation", "SL", "--absent", "Quinn")
    assert proposal["ok"]
    stuck = proposal["needs_placement"]
    assert [s["student"] for s in stuck] == ["Reyes, Noor"]
    assert "Brooks, Tyler would face two partners alone again" in stuck[0]["reason"]


def test_draft_moves_the_extra_student_so_nobody_is_the_one_twice():
    folder = workspace()
    # A fictional earlier simulation (written on the SL tab) where Grace Mahoney faced two partners alone.
    wb = load_workbook(folder / BEFORE)
    ws = wb["SL"]
    for r, (sup, sup_email, sub, sub_email) in enumerate([
            ("Mahoney, Grace", "gmahoney@example.edu", "Hale, Simone", "shale@example.edu"),
            (None, None, "Ibarra, Lucia", "libarra@example.edu")], start=2):
        ws.cell(r, 1).value = 1 if r == 2 else None
        if sup:
            ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 5).value = sup, sup_email, "SL_1"
        ws.cell(r, 6).value, ws.cell(r, 7).value, ws.cell(r, 9).value = sub, sub_email, "SL_1"
    wb.save(folder / BEFORE)
    result = run(folder, "draft", "--roster", BEFORE)
    assert result["ok"] and not result["errors"], result
    # Normally Tess Zimmerman (the extra Seller) joins Grace Mahoney's match 19; now she joins match 18.
    draft = {m["match"]: m for m in result["draft"]}
    assert draft[19]["SELLER"] == ["Yilmaz, Deniz"]
    assert draft[18]["SELLER"] == ["Walsh, Connor", "Zimmerman, Tess"]
    assert any("Mahoney, Grace" in n for n in result["placement_notes"])
    assert not any("faces two partners alone" in w for w in result["warnings"])


def test_salary_draft_rotates_the_right_column():
    folder = workspace()
    result = run(folder, "draft", "--roster", AFTER, "--simulation", "SL")
    assert result["ok"] and not result["errors"] and result["matches"] == 19, result
    draft = {m["match"]: (m["SUPERVISOR"], m["SUBORDINATE"]) for m in result["draft"]}
    uc = {m["match"]: m for m in EXPECTED["draft"]}
    buyers = [n for m in EXPECTED["draft"] for n in m["BUYER"]]
    sellers = [n for m in EXPECTED["draft"] for n in m["SELLER"]]
    # Used Car Buyers stay on the left in order; the right column moves down one and the bottom goes to the top.
    assert [draft[m][0] for m in sorted(draft)] == [[b] for b in buyers]
    assert draft[1][1] == [sellers[-1]] and draft[2][1] == [sellers[0]] and draft[3][1] == [sellers[1]]
    assert len(uc) == 19 and sellers[-1] == "Zimmerman, Tess"
    # The extra Subordinate would meet her Used Car partner in the last match, so she joins match 18 instead.
    assert draft[19][1] == ["Walsh, Connor"] and draft[18][1] == ["Van der Berg, Elise", "Yilmaz, Deniz"]
    assert any("already partners with Mahoney, Grace" in n for n in result["placement_notes"])
    # Rafael Mendes joined Maya Alvarez's match in class, so the rotation repeats that pair: shown, not fixed.
    assert result["warnings"] == ["Repeat partners in match 2: Alvarez, Maya and Mendes, Rafael were already "
                                  "together (UC_2)."]
    written = run(folder, "write-draft", "--roster", AFTER, "--simulation", "SL", "--accept-warnings")
    assert written["ok"], written
    tab = ns.read_sim_tab(load_workbook(folder / AFTER), "SL", ["SUPERVISOR", "SUBORDINATE"])
    assert len(tab.groups(tab.draft_members)) == 19


def test_salary_draft_keeps_the_sides_even_after_a_drop():
    folder = workspace()
    wb = load_workbook(folder / AFTER)
    ws = wb["Att - Fall"]
    for name in ("Adebayo, Tolu", "Alvarez, Maya"):  # two Used Car Buyers drop: 17 left, 20 right
        ws.delete_rows(next(c.row for c in ws["B"] if c.value == name))
    wb.save(folder / AFTER)
    result = run(folder, "draft", "--roster", AFTER, "--simulation", "SL")
    assert result["ok"] and result["students"] == 37
    left = [n for m in result["draft"] for n in m["SUPERVISOR"]]
    assert len(left) == 18 and "Zimmerman, Tess" in left
    assert any("Zimmerman, Tess moves to the Supervisor side" in n for n in result["placement_notes"])
    before = run(folder, "draft", "--roster", BEFORE, "--simulation", "SL")
    assert not before["ok"] and before["error"] == "previous_draft_missing"


BASELINE_CSV = "BUS+4489+F26+Week+1+-+Negotiation+Baseline+Survey_September+30,+2026_09.44.csv"
BASELINE = json.loads((EXAMPLES / "expected" / "baseline.json").read_text(encoding="utf-8"))
EVANS = "R_base0009=Evans, Jordan"   # "Jordan Evens" in the survey


def baseline_workspace() -> Path:
    folder = workspace()
    shutil.copy(EXAMPLES / BASELINE_CSV, folder / BASELINE_CSV)
    return folder


def test_baseline_preview_reads_gender_and_consent():
    folder = baseline_workspace()
    result = run(folder, "baseline", "--roster", DRAFTED, "--csv", BASELINE_CSV)
    assert result["ok"] and result["students"] == 39 and result["responses"] == 38   # preview row left out
    assert [u["typed_name"] for u in result["unresolved"]] == ["Jordan Evens"]
    assert result["unresolved"][0]["suggestions"] == ["Evans, Jordan"]
    result = run(folder, "baseline", "--roster", DRAFTED, "--csv", BASELINE_CSV, "--assign", EVANS)
    table = {row["student"]: [row["gender"], row["consent"]] for row in result["table"]}
    assert table == BASELINE["students"]
    assert result["declined"] == ["Bennett, Owen", "Ortiz, Lena"] and result["excluded"] == ["Hoang, Minh"]
    assert result["no_survey"] == ["Underwood, Mia", "Walsh, Connor"]
    assert result["duplicates"] == ["Sato, Hiro took the survey 2 times; the latest is used."]


def test_write_baseline_adds_two_columns_after_email():
    folder = baseline_workspace()
    others = all_values(folder / DRAFTED, skip_tab="Att - Fall")
    att_before = load_workbook(folder / DRAFTED)["Att - Fall"]
    refused = run(folder, "write-baseline", "--roster", DRAFTED, "--csv", BASELINE_CSV)
    assert not refused["ok"] and refused["error"] == "unresolved_responses"
    result = run(folder, "write-baseline", "--roster", DRAFTED, "--csv", BASELINE_CSV, "--assign", EVANS)
    assert result["ok"], result
    assert Path(result["backup"]).exists() and "inserted after Email" in result["written_to"]
    assert all_values(folder / DRAFTED, skip_tab="Att - Fall") == others, "a tab other than the class list changed"
    att = load_workbook(folder / DRAFTED)["Att - Fall"]
    assert [att.cell(2, c).value for c in range(1, 7)] == ["No.", "Student Name", "Email", "Gender", "Consent", "Goals"]
    for r in range(1, att_before.max_row + 1):   # everything right of Email moved over by two
        for c in range(1, att_before.max_column + 1):
            assert att.cell(r, c if c <= 3 else c + 2).value == att_before.cell(r, c).value, (r, c)
    rows = {att.cell(r, 2).value: r for r in range(3, att.max_row + 1)}
    for name, (gender, consent) in BASELINE["students"].items():
        assert [att.cell(rows[name], 4).value or "", att.cell(rows[name], 5).value] == [gender, consent], name
    # The other skills still read the class list, and a second run changes nothing.
    assert run(folder, "check", "--roster", DRAFTED)["ok"]
    again = run(folder, "write-baseline", "--roster", DRAFTED, "--csv", BASELINE_CSV, "--assign", EVANS)
    assert again["ok"] and "inserted" not in again["written_to"]
    # A different answer for someone already filled in needs her OK.
    changed = run(folder, "write-baseline", "--roster", DRAFTED, "--csv", BASELINE_CSV)
    assert not changed["ok"] and changed["error"] == "unresolved_responses"
    att2 = load_workbook(folder / DRAFTED)
    att2["Att - Fall"].cell(rows["Alvarez, Maya"], 5).value = "No"
    att2.save(folder / DRAFTED)
    differs = run(folder, "write-baseline", "--roster", DRAFTED, "--csv", BASELINE_CSV, "--assign", EVANS)
    assert not differs["ok"] and differs["error"] == "baseline_exists" and "Alvarez, Maya" in differs["changes"][0]


def test_write_baseline_asks_her_to_insert_columns_when_the_tab_has_formulas():
    folder = baseline_workspace()
    wb = load_workbook(folder / DRAFTED)
    att = wb["Att - Fall"]
    att.cell(45, 9).value = '=COUNTIF(I3:I41,"x")'   # a total under the attendance column
    wb.save(folder / DRAFTED)
    result = run(folder, "write-baseline", "--roster", DRAFTED, "--csv", BASELINE_CSV, "--assign", EVANS)
    assert not result["ok"] and result["error"] == "insert_columns_needed"
    assert any("formulas" in r for r in result["reasons"])
    # She inserts the two blank columns herself in Excel (simulated here, formula adjusted as Excel would).
    wb = load_workbook(folder / DRAFTED)
    att = wb["Att - Fall"]
    att.insert_cols(4, 2)
    att.cell(2, 4).value, att.cell(2, 5).value = "Gender", "Consent"
    att.cell(45, 11).value = '=COUNTIF(K3:K41,"x")'
    wb.save(folder / DRAFTED)
    filled = run(folder, "write-baseline", "--roster", DRAFTED, "--csv", BASELINE_CSV, "--assign", EVANS)
    assert filled["ok"] and "inserted" not in filled["written_to"], filled
    att = load_workbook(folder / DRAFTED)["Att - Fall"]
    assert att.cell(45, 11).value == '=COUNTIF(K3:K41,"x")' and att.cell(3, 5).value == "Yes"


def test_roles_are_locked_after_emails_go_out():
    folder = workspace()
    assert run(folder, "roles-sent", "--roster", DRAFTED)["ok"]
    result = run(folder, "write-draft", "--roster", DRAFTED, "--replace", "--swap", "Alvarez, Maya;Nakamura, Ren")
    assert not result["ok"] and result["error"] == "roles_locked"
    assert set(result["role_changes"]) == {"Alvarez, Maya", "Nakamura, Ren"}


def test_repeat_partners_are_warned_for_the_next_simulation():
    folder = workspace()
    wb = load_workbook(folder / AFTER)
    ws = wb["SL"]
    tab = ns.read_sim_tab(wb, "SL", ["SUPERVISOR", "SUBORDINATE"])
    r = tab.draft_header + 1
    ws.cell(r, 2).value, ws.cell(r, 3).value = "Alvarez, Maya", "Nakamura, Ren"      # partners in UC_2
    ws.cell(r + 1, 2).value, ws.cell(r + 1, 3).value = "Bennett, Owen", "Park, Daniel"  # never partners
    wb.save(folder / AFTER)
    result = run(folder, "check", "--roster", AFTER, "--simulation", "SL")
    assert result["ok"]
    repeats = [w for w in result["warnings"] if w.startswith("Repeat partners")]
    assert len(repeats) == 1 and "UC_2" in repeats[0] and "Alvarez, Maya" in repeats[0]


if __name__ == "__main__":
    failures = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            try:
                fn()
                print(f"PASS {name}")
            except AssertionError as exc:
                failures += 1
                print(f"FAIL {name}: {exc}")
    sys.exit(1 if failures else 0)

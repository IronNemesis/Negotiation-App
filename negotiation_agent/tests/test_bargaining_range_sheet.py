"""Tests for the bargaining-range-sheet skill, run against the fictional example files.

    python negotiation_agent/tests/test_bargaining_range_sheet.py

Works with pytest too. Each test copies the examples into a temporary folder, so the repo is never changed.
"""
from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = ROOT / "examples"
SCRIPT = ROOT / "skills" / "bargaining-range-sheet" / "scripts" / "brs.py"
CSV = "BUS+4489+F26+Week+2+-+Used+Car+Bargaining+Range_October+14,+2026_14.40.csv"
DRAFTED = "BUS 4489 F26 Roster - Used Car drafted.xlsx"
AFTER = "BUS 4489 F26 Roster - after Used Car.xlsx"
SHEET = "BUS 4489 F26 Week 2 - Used Car Bargaining Range.xlsx"
DECISIONS = ["--latest-duplicates", "--assign", "R_example0007=Evans, Jordan",
             "--assign", "R_example0026=Park, Danielle"]
EXPECTED = json.loads((EXAMPLES / "expected" / "used-car.json").read_text(encoding="utf-8"))
ORANGE, LIGHT_RED, YELLOW = "F8CBAD", "FFC7CE", "FFFF00"


def run(folder: Path, *args) -> dict:
    out = subprocess.run([sys.executable, str(SCRIPT), *args], cwd=folder, capture_output=True, text=True)
    assert out.stdout, out.stderr
    return json.loads(out.stdout)


def workspace() -> Path:
    tmp = Path(tempfile.mkdtemp())
    for f in (DRAFTED, AFTER, CSV):
        shutil.copy(EXAMPLES / f, tmp / f)
    return tmp


def build(folder: Path, roster: str, *extra) -> dict:
    return run(folder, "build", "--roster", roster, "--csv", CSV, "--week", "2", *DECISIONS, *extra)


def sheet_rows(path: Path):
    ws = load_workbook(path).worksheets[0]
    header = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 2).value is None and ws.cell(r, header.index("OUTCOME") + 3).value is None:
            break
        rows.append(r)
    return ws, header, rows


def people_on_sheet(ws, header, rows) -> dict:
    """{'Last, First': (row, side)} for every name on the sheet."""
    first_cols = [i + 1 for i, h in enumerate(header) if h == "FIRST"]
    out = {}
    for r in rows:
        for side, c in zip(("BUYER", "SELLER"), first_cols):
            first, last = ws.cell(r, c).value, ws.cell(r, c + 1 if side == "BUYER" else c + 1).value
            if first:
                out[f"{last}, {first}"] = (r, side, c)
    return out


def fill_of(cell) -> str | None:
    return cell.fill.fgColor.rgb[-6:] if cell.fill and cell.fill.fill_type == "solid" else None


def test_cleaning_rules():
    sys.path.insert(0, str(SCRIPT.parent))
    import brs
    cases = {
        "7800": (7800, ""), "$7800": (7800, ""), "7,800": (7800, ""), "$7,800 ": (7800, ""), "$13000": (13000, ""),
        "10": (10000, "corrected"), "9.4": (9400, "corrected"), "8.8k": (8800, "corrected"),
        "8.8K": (8800, "corrected"), "$8,800 (sell to dealer)": (8800, "corrected"),
        "8800 (dealer)": (8800, "corrected"),
        ">8,800": (">8,800", "unclear"), "$10,500-10,250": ("$10,500-10,250", "unclear"),
        "(8800)": ("(8800)", "unclear"), "about 9000": ("about 9000", "unclear"),
        "200": (200, "out_of_range"), "0": (0, "out_of_range"), "": (None, ""),
    }
    for typed, (value, flag) in cases.items():
        answer = brs.clean_price(typed, 1000, 50000)
        assert (answer.value, answer.flag) == (value, flag), (typed, answer)


def test_prepare_finds_every_planted_problem():
    folder = workspace()
    result = run(folder, "prepare", "--roster", DRAFTED, "--csv", CSV)
    assert result["ok"] and result["matches_from"] == "draft table" and result["matches"] == 19
    assert {r["reason"].split(" ")[0] for r in result["responses_removed"]} == {"every", "not"}
    unresolved = {u.get("typed_name") or u.get("student"): u for u in result["unresolved"]}
    assert set(unresolved) == {"Jordan Evens", "Dani Park", "Hale, Simone"}
    assert unresolved["Jordan Evens"]["suggestions"] == ["Evans, Jordan"]
    assert set(unresolved["Dani Park"]["suggestions"]) == {"Park, Daniel", "Park, Danielle"}
    assert result["matched"].get("nickname") == 2
    assert result["role_conflicts"] == ["Shah, Dev: Chose Buyer in the survey; assigned Seller."]
    cleaning = {(c["student"], c["question"]): k for k, items in result["cleaning"].items() for c in items}
    assert cleaning == {("Reyes, Noor", "TARGET"): "corrected", ("Walsh, Connor", "BATNA"): "corrected",
                        ("Silva, Kai", "RESERVATION"): "unclear", ("Rossi, Gianna", "TARGET"): "unclear",
                        ("Thornton, Jack", "BATNA"): "corrected", ("Sato, Hiro", "BATNA"): "out_of_range"}
    assert not result["ready_to_build"]


def test_build_refuses_until_decisions_are_made():
    folder = workspace()
    result = run(folder, "build", "--roster", DRAFTED, "--csv", CSV, "--week", "2")
    assert not result["ok"] and result["error"] == "unresolved_responses"
    assert not (folder / SHEET).exists()


def test_sheet_from_draft_matches_her_layout():
    folder = workspace()
    roster_hash = hashlib.sha256((folder / DRAFTED).read_bytes()).hexdigest()
    result = build(folder, DRAFTED)
    assert result["ok"], result
    assert hashlib.sha256((folder / DRAFTED).read_bytes()).hexdigest() == roster_hash, "roster workbook changed"
    ws, header, rows = sheet_rows(folder / SHEET)
    assert header[:9] == ["MATCH", "FIRST", "LAST", "ROLE", "ACTUAL IO", "INITIAL", "TARGET", "RESERVATION", "BATNA"]
    outcome = header.index("OUTCOME")
    assert header[outcome + 1:outcome + 9] == ["BATNA", "RESERVATION", "TARGET", "INITIAL", "ACTUAL IO", "FIRST",
                                               "LAST", "ROLE"]
    assert all(fill_of(ws.cell(1, i + 1)) == YELLOW for i, h in enumerate(header) if h == "ACTUAL IO")
    # One row per match, plus Tess Zimmerman's extra row under match 19.
    matches = [ws.cell(r, 1).value for r in rows]
    assert [m for m in matches if m is not None] == list(range(1, 20))
    assert matches[-1] is None and len(rows) == 20
    people = people_on_sheet(ws, header, rows)
    assert len(people) == 39
    assert people["Zimmerman, Tess"][0] == rows[-1] and people["Zimmerman, Tess"][1] == "SELLER"
    # Every student is on the side of their drafted role, in their drafted match.
    for m in EXPECTED["draft"]:
        for role in ("BUYER", "SELLER"):
            for name in m[role]:
                assert people[name][1] == role, name
    # In-class columns are empty.
    for i, h in enumerate(header):
        if h in ("ACTUAL IO", "OUTCOME"):
            assert all(ws.cell(r, i + 1).value is None for r in rows), h
    # Planted answers.
    def cell(name, question):
        r, side, first_col = people[name]
        offset = {"BUYER": {"INITIAL": 4, "TARGET": 5, "RESERVATION": 6, "BATNA": 7, "ROLE": 2},
                  "SELLER": {"INITIAL": -2, "TARGET": -3, "RESERVATION": -4, "BATNA": -5, "ROLE": 2}}[side]
        return ws.cell(r, first_col + offset[question])
    assert cell("Reyes, Noor", "TARGET").value == 10000 and fill_of(cell("Reyes, Noor", "TARGET")) == ORANGE
    assert "10" in cell("Reyes, Noor", "TARGET").comment.text
    assert cell("Walsh, Connor", "BATNA").value == 8800 and fill_of(cell("Walsh, Connor", "BATNA")) == ORANGE
    assert cell("Silva, Kai", "RESERVATION").value == ">8,800"
    assert fill_of(cell("Silva, Kai", "RESERVATION")) == LIGHT_RED
    assert cell("Thornton, Jack", "BATNA").value == 8800 and fill_of(cell("Thornton, Jack", "BATNA")) == ORANGE
    assert "sell to dealer" in cell("Thornton, Jack", "BATNA").comment.text
    assert cell("Sato, Hiro", "BATNA").value == 200 and fill_of(cell("Sato, Hiro", "BATNA")) == LIGHT_RED
    assert cell("Chen, Priya", "INITIAL").value == 7000 and fill_of(cell("Chen, Priya", "INITIAL")) != ORANGE
    assert cell("Shah, Dev", "ROLE").value == 2 and fill_of(cell("Shah, Dev", "ROLE")) == LIGHT_RED
    assert cell("Bennett, Owen", "INITIAL").value is None
    assert isinstance(cell("Evans, Jordan", "INITIAL").value, int)
    # Every price that isn't flagged is a plain number.
    for name in people:
        for q in ("INITIAL", "TARGET", "RESERVATION", "BATNA"):
            c = cell(name, q)
            if c.value is not None and fill_of(c) not in (ORANGE, LIGHT_RED):
                assert isinstance(c.value, int), (name, q, c.value)
    wb = load_workbook(folder / SHEET)
    assert wb.sheetnames[1:] == ["Roles", "Raw"]
    assert wb["Raw"].max_row == len((folder / CSV).read_text(encoding="utf-8-sig").splitlines())


def test_rerun_after_absences_uses_final_matches():
    folder = workspace()
    assert build(folder, DRAFTED)["ok"]
    result = build(folder, AFTER)
    assert result["ok"], result
    assert result["matches_from"] == "final table" and result["matches"] == 17
    assert result["previous_version_backup"] and Path(result["previous_version_backup"]).exists()
    assert any(s.startswith("Castellano, Rosa (absent") for s in result["left_off_sheet"])
    ws, header, rows = sheet_rows(folder / SHEET)
    people = people_on_sheet(ws, header, rows)
    assert set(people) == {n for m in EXPECTED["final"] for r in ("BUYER", "SELLER") for n in m[r]}
    assert not {"Adebayo, Tolu", "Castellano, Rosa", "Zimmerman, Tess"} & set(people)
    # Rafael Mendes sits on the row under match 2; Siobhan O'Neill under match 4.
    for name, match in (("Mendes, Rafael", 2), ("O'Neill, Siobhan", 4)):
        r = people[name][0]
        assert ws.cell(r, 1).value is None and ws.cell(r - 1, 1).value == match, name


def test_existing_in_class_entries_are_protected():
    folder = workspace()
    assert build(folder, DRAFTED)["ok"]
    wb = load_workbook(folder / SHEET)
    ws = wb.worksheets[0]
    outcome_col = [c.value for c in ws[1]].index("OUTCOME") + 1
    ws.cell(2, outcome_col, "9.2K")
    wb.save(folder / SHEET)
    result = build(folder, AFTER)
    assert not result["ok"] and result["error"] == "in_class_entries_present"
    assert load_workbook(folder / SHEET).worksheets[0].cell(2, outcome_col).value == "9.2K"


if __name__ == "__main__":
    failures = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            try:
                fn()
                print(f"PASS {name}")
            except AssertionError as exc:
                failures += 1
                print(f"FAIL {name}: {exc}")
    sys.exit(1 if failures else 0)

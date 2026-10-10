"""Generate the fictional example files in the instructor's exact formats.

Run from the repo root:

    python negotiation_agent/examples/tools/make_examples.py

Writes, under negotiation_agent/examples/:
  BUS 4489 F26 Roster - before Used Car.xlsx   UC and SL tabs prepared but empty
  BUS 4489 F26 Roster - after Used Car.xlsx    UC draft and final filled in (with absences), SL prepared
  BUS+4489+F26+Week+2+-+Used+Car+Bargaining+Range_October+14,+2026_14.40.csv
  expected/used-car.json                        the results each skill should produce

Every name, email, and answer is fictional. Emails use example.edu (class list)
and example.com (personal), so nothing can reach a real person. The contents are
deterministic: running it twice gives the same names, matches, and answers.
"""
from __future__ import annotations

import csv
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent
EXAMPLES = HERE.parent
sys.path.insert(0, str(EXAMPLES.parent / "shared"))
from nps_shared import name_key, split_class_name  # noqa: E402

COURSE = "BUS 4489 F26"
DATE = "2026-10-14"

CLASS = [
    "Adebayo, Tolu", "Alvarez, Maya", "Bennett, Owen", "Brooks, Tyler", "Castellano, Rosa", "Chen, Priya",
    "Dawson, Leah", "Delgado, Marco", "Evans, Jordan", "Fitzgerald, Nora", "García, Mateo", "Hale, Simone",
    "Hoang, Minh", "Ibarra, Lucia", "Jensen, Kai", "Kowalski, Andy", "Laurent, Camille", "Lindqvist, Erik",
    "Mahoney, Grace", "Mendes, Rafael", "Nakamura, Ren", "Nguyen, Katherine", "Novak, Ben", "O'Neill, Siobhan",
    "Ortiz, Lena", "Park, Daniel", "Park, Danielle", "Quinn, Ava", "Reyes, Noor", "Rossi, Gianna", "Sato, Hiro",
    "Shah, Dev", "Silva, Kai", "Thornton, Jack", "Underwood, Mia", "Van der Berg, Elise", "Walsh, Connor",
    "Yilmaz, Deniz", "Zimmerman, Tess",
]
ABSENT_IN_CLASS = ["Adebayo, Tolu", "Castellano, Rosa", "Zimmerman, Tess"]


def _local_part(name: str, first_letters: int) -> str:
    first, last = split_class_name(name)
    local = (first[:first_letters] + "".join(ch for ch in last if ch.isalpha())).lower()[:8]
    return local.replace("í", "i").replace("á", "a")


def _build_emails() -> dict[str, str]:
    """Cal Poly style addresses (first initial + last name, 8 letters), made unique like a real campus would."""
    emails, taken = {}, set()
    for name in sorted(CLASS, key=lambda n: n):
        letters = 1
        while (local := _local_part(name, letters)) in taken:
            letters += 1
        taken.add(local)
        emails[name] = f"{local}@example.edu"
    return emails


def email_for(name: str) -> str:
    return EMAILS[name]


EMAILS: dict[str, str] = {}


def sort_key(name: str) -> str:
    first, last = split_class_name(name)
    return name_key(first, last)


# ---------------------------------------------------------------- expected pairing results

def draft_matches(names: list[str]) -> list[dict]:
    """Her Used Car method: alphabetical, first half Buyers, second half Sellers, matched in order."""
    ordered = sorted(names, key=sort_key)
    half = len(ordered) // 2
    buyers, sellers = ordered[:half], ordered[half:]
    matches = [{"match": i + 1, "BUYER": [b], "SELLER": [s]} for i, (b, s) in enumerate(zip(buyers, sellers))]
    if len(sellers) > len(buyers):
        matches[-1]["SELLER"].append(sellers[-1])
    return matches


def apply_absences(draft: list[dict], absent: list[str]) -> list[dict]:
    """Her absence rule: stranded students join the match just before theirs (or the next, for match 1)."""
    matches = [{"match": m["match"], "BUYER": list(m["BUYER"]), "SELLER": list(m["SELLER"])} for m in draft]
    for m in matches:
        for role in ("BUYER", "SELLER"):
            m[role] = [n for n in m[role] if n not in absent]
    for i, m in enumerate(matches):
        if m["BUYER"] and m["SELLER"]:
            continue
        stranded = [(role, n) for role in ("BUYER", "SELLER") for n in m[role]]
        m["BUYER"], m["SELLER"] = [], []
        alive = [x for x in matches if x["BUYER"] or x["SELLER"]]
        earlier = [x for x in alive if x["match"] < m["match"]]
        target = earlier[-1] if earlier else next(x for x in alive if x["match"] > m["match"])
        for role, n in stranded:
            target[role].append(n)
    return [m for m in matches if m["BUYER"] or m["SELLER"]]


# ---------------------------------------------------------------- workbooks

ROLE_HEADER = ["BUYER", "Email", None, "Gender", "SELLER", "Email", None, "Gender", "Intervention"]


def add_attendance(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Att - Fall"
    weeks = ["Aug. 25", "Sep. 1", "Sept. 8", "Sep. 15", "Sept. 22", "Sep. 29", "Oct. 6", "Oct. 13"]
    ws.append([None, None, None, None, "Baseline", None] + [f"Week {i}" for i in range(1, len(weeks) + 1)])
    ws.append(["No.", "Student Name", "Email", "Goals", "Part I", "Part II"] + weeks)
    for i, name in enumerate(sorted(CLASS, key=sort_key), start=1):
        ws.append([i, name, email_for(name), None, "x", "x"] + ["x"] * (len(weeks) - 1))
    for c in ("B", "C"):
        ws.column_dimensions[c].width = 24
    spring = wb.create_sheet("Att - Spring")
    spring.append([None, None, None, None, "Baseline", None, None, "Week 1", None, "Week 2"])
    spring.append(["No.", "Student Name", "Email", "Credit/NC", "Goals", "Part I", "Part II", "Jan. 9", "Jan. 11", "Jan. 16"])
    for i in range(1, 41):
        spring.append([i])


def add_exam(wb: Workbook) -> None:
    ws = wb.create_sheet("E1")
    ws.append(["Student ID", "Team ID", "Student", "Group SA", "Group MC", "Group Total", "Group Curved"])
    for i, name in enumerate(sorted(CLASS, key=sort_key)[:10], start=1):
        r = i + 1
        ws.append([i, (i - 1) // 5 + 1, name, 20 + i % 5, 18 + i % 4, f"=SUM(D{r}:E{r})", f"=F{r}+1.5"])
    ws["G1"].comment = Comment("Curve agreed with the department chair.", "Instructor")


def add_prepared_tab(wb: Workbook, code: str, role_a: str, role_b: str, prior: list[str]) -> None:
    """A tab prepared the way she prepares them: headers and match numbers, no names."""
    ws = wb.create_sheet(code)
    ids = prior + [code]
    if code == "UC":
        ws.append([None] + ROLE_HEADER)
    else:
        ws.append(["MATCH", role_a, "Email", *ids, role_b, "Email", *ids])
    for i in range(1, 23):
        ws.cell(i + 1, 1, i)
    ws.cell(29, 1, "MATCH")
    ws.cell(29, 2, role_a)
    ws.cell(29, 3, role_b)
    for i in range(1, 23):
        ws.cell(29 + i, 1, i)
    for c in "BF":
        ws.column_dimensions[c].width = 24
    for c in "CG":
        ws.column_dimensions[c].width = 22


def fill_used_car(ws, draft: list[dict], final: list[dict], absent: list[str]) -> None:
    """Fill the UC tab the way the skills should: draft table at the bottom, final table at the top."""
    r = 30
    for m in draft:
        ws.cell(r, 1, m["match"])
        ws.cell(r, 2, m["BUYER"][0])
        ws.cell(r, 3, m["SELLER"][0])
        r += 1
        for extra in m["SELLER"][1:]:
            ws.cell(r, 1).value = None
            ws.cell(r, 3, extra)
            r += 1
    for row in range(r, 52):
        if isinstance(ws.cell(row, 1).value, int):
            ws.cell(row, 1).value = None
    r = 2
    for m in final:
        code = f"UC_{m['match']}"
        rows = max(len(m["BUYER"]), len(m["SELLER"]))
        for k in range(rows):
            ws.cell(r, 1).value = m["match"] if k == 0 else None
            for role, col in (("BUYER", 2), ("SELLER", 6)):
                if k < len(m[role]):
                    name = m[role][k]
                    ws.cell(r, col, name)
                    ws.cell(r, col + 1, email_for(name))
                    ws.cell(r, col + 2, code)
                else:
                    for off in range(3):
                        ws.cell(r, col + off).value = None
            r += 1
    for row in range(r, 24):
        ws.cell(row, 1).value = None
    r += 1
    for name in absent:
        ws.cell(r, 2, name)
        ws.cell(r, 3, email_for(name))
        ws.cell(r, 4, "UC_0")
        r += 1


def build_workbook(after: bool, draft, final) -> Workbook:
    wb = Workbook()
    add_attendance(wb)
    add_exam(wb)
    add_prepared_tab(wb, "UC", "BUYER", "SELLER", [])
    add_prepared_tab(wb, "SL", "SUPERVISOR", "SUBORDINATE", ["UC"])
    if after:
        fill_used_car(wb["UC"], draft, final, ABSENT_IN_CLASS)
    for ws in wb.worksheets:
        for c in ws[1]:
            if c.value:
                c.font = Font(bold=True)
    return wb


# ---------------------------------------------------------------- Qualtrics export

ITEM_TEXT = "After this negotiation is over, to what extent do you believe that your partner will: - "
ITEMS = ["Not be interested in working with you again.", "Not be interested in interacting socially with you.",
         "Not want to have you on their team if you worked together.",
         "Not be interested in going out for drinks with you after work.",
         "Not see you as the type of person they would want to work with.",
         "Not see you as the type of person they would want to socialize with."]
CODES = ["StartDate", "EndDate", "Status", "IPAddress", "Progress", "Duration (in seconds)", "Finished",
         "RecordedDate", "ResponseId", "RecipientLastName", "RecipientFirstName", "RecipientEmail",
         "ExternalReference", "LocationLatitude", "LocationLongitude", "DistributionChannel", "UserLanguage",
         "Last Seen Flow Element ID", "Last Seen Question IDs", "Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7",
         "Q12_1", "Q12_2", "Q12_3", "Q12_4", "Q12_5", "Q12_6", "Q11"]
QUESTIONS = ["Start Date", "End Date", "Response Type", "IP Address", "Progress", "Duration (in seconds)",
             "Finished", "Recorded Date", "Response ID", "Recipient Last Name", "Recipient First Name",
             "Recipient Email", "External Data Reference", "Location Latitude", "Location Longitude",
             "Distribution Channel", "User Language", "Last Seen Flow Element ID", "Last Seen Question IDs",
             "What is your FIRST NAME?", "What is your LAST NAME?", "What is your ROLE?",
             "1. What is the price you will initially offer the other party? (initial offer)",
             "2. What is the IDEAL price you would like to buy/sell the car for? (target price)",
             "3. What is the absolute highest price you will pay for the car if you are the buyer? OR What is the "
             "absolute lowest price you will sell the car for if you are the seller? (reservation price)",
             "4. What is the cost in a dollar amount of your best alternative (BATNA) if you are not able to come "
             "to an agreement in this negotiation? (need to give this in a dollar amount)",
             *[ITEM_TEXT + t for t in ITEMS],
             "5. What is your email address? (This will be used to immediately trigger an email with a summary of "
             "your responses to this survey. The email summary will come from instructor@example.edu. If you "
             "cannot find it, please look in your clutter or junk mail folder.)"]
IMPORT_IDS = ['{"ImportId":"startDate","timeZone":"America/Los_Angeles"}',
              '{"ImportId":"endDate","timeZone":"America/Los_Angeles"}', '{"ImportId":"status"}',
              '{"ImportId":"ipAddress"}', '{"ImportId":"progress"}', '{"ImportId":"duration"}',
              '{"ImportId":"finished"}', '{"ImportId":"recordedDate","timeZone":"America/Los_Angeles"}',
              '{"ImportId":"_recordId"}', '{"ImportId":"recipientLastName"}', '{"ImportId":"recipientFirstName"}',
              '{"ImportId":"recipientEmail"}', '{"ImportId":"externalDataReference"}',
              '{"ImportId":"locationLatitude"}', '{"ImportId":"locationLongitude"}',
              '{"ImportId":"distributionChannel"}', '{"ImportId":"userLanguage"}', '{"ImportId":"LastSeenFlowID"}',
              '{"ImportId":"LastSeenQuestions"}', '{"ImportId":"QID1_TEXT"}', '{"ImportId":"QID2_TEXT"}',
              '{"ImportId":"QID3"}', '{"ImportId":"QID4_TEXT"}', '{"ImportId":"QID5_TEXT"}',
              '{"ImportId":"QID6_TEXT"}', '{"ImportId":"QID7_TEXT"}',
              *[f'{{"ImportId":"QID12_{i}"}}' for i in range(1, 7)], '{"ImportId":"QID11_TEXT"}']

# Planted cases: student -> what is special about their response. Prices are (initial, target, reservation, BATNA).
SPECIAL = {
    "Reyes, Noor": {"prices": ("12000", "10", "9000", "8800"), "expect": {"TARGET": ("corrected", 10000)}},
    "Walsh, Connor": {"prices": ("11500", "10250", "9000", "8.8k"), "expect": {"BATNA": ("corrected", 8800)}},
    "Silva, Kai": {"prices": ("11,900", "10,600", ">8,800", "8,800"), "expect": {"RESERVATION": ("unclear", ">8,800")}},
    "Rossi, Gianna": {"prices": ("$11,000", "$10,500-10,250", "$9,400", "$8,880"), "first": "Gianna ", "last": "Rossi ",
                      "expect": {"TARGET": ("unclear", "$10,500-10,250")}},
    "Thornton, Jack": {"prices": ("$11,000", "$9,500", "$9,000", "$8,800 (sell to dealer)"),
                       "expect": {"BATNA": ("unclear", "$8,800 (sell to dealer)")}},
    "Sato, Hiro": {"prices": ("14000", "9600", "9000", "200"), "expect": {"BATNA": ("out_of_range", 200)}},
    "Ortiz, Lena": {"first": "lena", "last": "ortiz", "match": "exact"},
    "Kowalski, Andy": {"first": "Andrew", "email": "andrew.kowalski@example.com", "match": "nickname"},
    "Nguyen, Katherine": {"first": "Katie", "email": "katienguyen@example.com", "match": "nickname"},
    "Park, Danielle": {"first": "Dani", "email": "danipark@example.com", "match": "ask (two Parks)"},
    "García, Mateo": {"last": "Garcia", "match": "exact (accent ignored)"},
    "O'Neill, Siobhan": {"last": "O’Neill", "match": "exact (curly apostrophe ignored)"},
    "Van der Berg, Elise": {"last": "van der berg", "match": "exact"},
    "Evans, Jordan": {"last": "Evens", "email": "jevens22@example.com", "match": "ask (misspelled last name)"},
    "Shah, Dev": {"role": "1", "match": "exact", "expect": {"ROLE": ("role_conflict", 2)}},
    "Hale, Simone": {"duplicate": True, "match": "exact (submitted twice)"},
    "Mahoney, Grace": {"email": "gracem@example.com", "match": "exact"},
    "Chen, Priya": {"prices": ("$7,000", "$8,000", "$9,000", "$7,600"), "match": "exact"},
}
NO_RESPONSE = ["Bennett, Owen", "Underwood, Mia", "Adebayo, Tolu"]


def fmt_money(rng: random.Random, value: int) -> str:
    return rng.choice([str(value), f"${value}", f"{value:,}", f"${value:,}", f"{value} "])


def prices(rng: random.Random, role: str) -> tuple[int, int, int, int]:
    if role == "BUYER":
        reservation = rng.choice(range(8000, 10300, 100))
        target = reservation - rng.choice(range(400, 1600, 50))
        initial = target - rng.choice(range(200, 1500, 50))
        return initial, target, reservation, rng.choice([7600, 7600, 7600, 8000, 8200])
    reservation = rng.choice(range(8800, 9600, 50))
    target = reservation + rng.choice(range(400, 2000, 50))
    initial = target + rng.choice(range(500, 3000, 50))
    return initial, target, reservation, rng.choice([8800, 8800, 8800, 8900, 9000])


def build_survey(draft: list[dict]) -> tuple[list[list[str]], dict]:
    rng = random.Random(4489)
    role_of = {n: role for m in draft for role in ("BUYER", "SELLER") for n in m[role]}
    rows, expected = [], {}
    minute = 0

    def stamp(extra: int = 0) -> tuple[str, str]:
        nonlocal minute
        minute += rng.randint(9, 41) + extra
        start = 15 * 60 + minute
        day, mins = 12 + start // (24 * 60), start % (24 * 60)
        end = mins + rng.randint(2, 15)
        return (f"2026-10-{day:02d} {mins // 60:02d}:{mins % 60:02d}:{rng.randint(0, 59):02d}",
                f"2026-10-{day:02d} {end // 60 % 24:02d}:{end % 60:02d}:{rng.randint(0, 59):02d}")

    def row(first, last, role_code, price_text, email, status="0", extra=0):
        start, end = stamp(extra)
        items = [str(rng.randint(1, 4)) for _ in ITEMS]
        n = len(rows) + 1
        return [start, end, status, f"192.0.2.{n}", "100", str(rng.randint(80, 3000)), "1", end, f"R_example{n:04d}",
                "", "", "", "", "", "", "anonymous", "EN", "FL_EOS_ID", "", first, last, role_code, *price_text,
                *items, email]

    for name in sorted(CLASS, key=sort_key):
        if name in NO_RESPONSE:
            expected[name] = {"response": "none"}
            continue
        spec = SPECIAL.get(name, {})
        first, last = split_class_name(name)
        role = role_of[name]
        role_code = spec.get("role", "1" if role == "BUYER" else "2")
        price_text = spec.get("prices") or tuple(fmt_money(rng, v) for v in prices(rng, role))
        email = spec.get("email", email_for(name) if rng.random() < 0.7 else f"{first.lower()}{rng.randint(10, 99)}@example.com")
        rows.append(row(spec.get("first", first), spec.get("last", last), role_code, price_text, email))
        expected[name] = {"match": spec.get("match", "exact or email"), "flags": spec.get("expect", {})}
        if spec.get("duplicate"):
            later = tuple(fmt_money(rng, v) for v in prices(rng, role))
            rows.append(row(first, last, role_code, later, email, extra=120))
    rows.insert(3, row("", "", "", ("", "", "", ""), "", extra=0))
    for c in range(19, 33):
        rows[3][c] = ""
    rows.insert(0, row("Test", "Preview", "1", ("7000", "8000", "9000", "7600"), "instructor@example.edu", status="1"))
    rows.sort(key=lambda r: r[7])
    return rows, expected


def main() -> None:
    EMAILS.update(_build_emails())
    draft = draft_matches(CLASS)
    final = apply_absences(draft, ABSENT_IN_CLASS)
    (EXAMPLES / "expected").mkdir(exist_ok=True)
    build_workbook(False, draft, final).save(EXAMPLES / f"{COURSE} Roster - before Used Car.xlsx")
    build_workbook(True, draft, final).save(EXAMPLES / f"{COURSE} Roster - after Used Car.xlsx")

    rows, survey_expected = build_survey(draft)
    csv_name = "BUS+4489+F26+Week+2+-+Used+Car+Bargaining+Range_October+14,+2026_14.40.csv"
    with (EXAMPLES / csv_name).open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(CODES)
        w.writerow(QUESTIONS)
        w.writerow(IMPORT_IDS)
        w.writerows(rows)

    expected = {
        "class_size": len(CLASS),
        "draft": draft,
        "absent_in_class": ABSENT_IN_CLASS,
        "final": final,
        "survey": {"rows_in_export": len(rows), "blank_rows": 1, "preview_rows": 1,
                   "students": survey_expected},
    }
    (EXAMPLES / "expected" / "used-car.json").write_text(json.dumps(expected, indent=2, ensure_ascii=False) + "\n",
                                                          encoding="utf-8")
    print(f"Draft: {len(draft)} matches; final after absences: {len(final)} matches; survey rows: {len(rows)}")


if __name__ == "__main__":
    main()

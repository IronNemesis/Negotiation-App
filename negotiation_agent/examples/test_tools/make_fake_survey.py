"""Create a fake Qualtrics export for testing Stage 2.

Reads the approved Pairing Draft.xlsx in a simulation folder and writes
Survey Export.csv next to it, in Qualtrics' three-header-row CSV format.
Planted exceptions (so H6 and T8's rules get exercised):
  - one "Survey Preview" response that should be ignored
  - one student who submitted twice
  - one response with a mistyped student ID (that student has no other response)
  - two students who never responded
  - student IDs with a leading zero are written without it, as Excel would

Usage:
    python make_fake_survey.py "<simulation folder>"
"""
import csv
import random
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

RANGES = {  # role -> (worst acceptable, ideal, opening offer)
    "buyer": ((13000, 15500), (10000, 12000), (9000, 11000)),
    "seller": ((11000, 13000), (15000, 17000), (16000, 18000)),
}


def find_header(rows):
    for i, row in enumerate(rows):
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if any("id" in c for c in cells) and any("role" in c for c in cells):
            return i, cells
    sys.exit("Could not find a header row with student ID and role columns in Pairing Draft.xlsx")


def column(cells, *words, avoid="partner"):
    for i, c in enumerate(cells):
        if all(w in c for w in words) and avoid not in c:
            return i
    sys.exit(f"Could not find a column containing {words} in Pairing Draft.xlsx")


def main():
    folder = Path(sys.argv[1])
    rows = list(load_workbook(folder / "Pairing Draft.xlsx", data_only=True).worksheets[0].iter_rows(values_only=True))
    h, cells = find_header(rows)
    id_col, role_col = column(cells, "id"), column(cells, "role")
    students = [(str(r[id_col]).strip(), str(r[role_col]).strip()) for r in rows[h + 1:] if r[id_col]]
    if len(students) < 6:
        sys.exit("Need at least 6 students in the draft to plant all the exceptions.")

    rng = random.Random(7)
    shuffled = students[:]
    rng.shuffle(shuffled)
    no_response = shuffled[:2]
    typo_student = shuffled[2]
    duplicate_student = shuffled[3]

    start = datetime(2026, 10, 14, 18, 0)
    responses = []

    def add(sid, role, status="IP Address", offset=None):
        key = "buyer" if role.lower().startswith("b") else "seller"
        vals = [rng.randrange(lo, hi, 100) for lo, hi in RANGES[key]]
        t = start + timedelta(minutes=offset if offset is not None else rng.randint(0, 600))
        responses.append((t, status, sid, role, vals))

    for sid, role in students:
        if (sid, role) in no_response:
            continue
        if (sid, role) == typo_student:
            add(sid[:-2] + sid[-1] + sid[-2], role)  # last two digits swapped
            continue
        add(sid.lstrip("0") if sid.startswith("0") else sid, role)
    add(duplicate_student[0], duplicate_student[1], offset=700)  # second, later submission
    add("123456789", "Buyer", status="Survey Preview", offset=-60)
    responses.sort(key=lambda r: r[0])

    out = folder / "Survey Export.csv"
    with out.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["StartDate", "EndDate", "Status", "Progress", "Duration (in seconds)", "Finished",
                    "RecordedDate", "ResponseId", "Q1", "Q2", "Q3", "Q4", "Q5"])
        w.writerow(["Start Date", "End Date", "Response Type", "Progress", "Duration (in seconds)", "Finished",
                    "Recorded Date", "Response ID", "What is your Cal Poly student ID?", "What is your role?",
                    "What is the worst price you would accept?", "What is your ideal price?",
                    "What will your opening offer be?"])
        w.writerow(['{"ImportId":"startDate"}', '{"ImportId":"endDate"}', '{"ImportId":"status"}',
                    '{"ImportId":"progress"}', '{"ImportId":"duration"}', '{"ImportId":"finished"}',
                    '{"ImportId":"recordedDate"}', '{"ImportId":"_recordId"}', '{"ImportId":"QID1_TEXT"}',
                    '{"ImportId":"QID2"}', '{"ImportId":"QID3_TEXT"}', '{"ImportId":"QID4_TEXT"}',
                    '{"ImportId":"QID5_TEXT"}'])
        for i, (t, status, sid, role, vals) in enumerate(responses, 1):
            end = t + timedelta(seconds=rng.randint(90, 400))
            w.writerow([t.strftime("%Y-%m-%d %H:%M:%S"), end.strftime("%Y-%m-%d %H:%M:%S"), status, 100,
                        int((end - t).total_seconds()), "True", end.strftime("%Y-%m-%d %H:%M:%S"),
                        f"R_fake{i:04d}", sid, role, *vals])

    print(f"Wrote {out}")
    print("Planted exceptions (what the agent should catch):")
    print(f"  No response:        {no_response[0][0]}, {no_response[1][0]}")
    print(f"  Mistyped ID:        real ID {typo_student[0]} was entered with its last two digits swapped")
    print(f"  Duplicate:          {duplicate_student[0]} submitted twice (second one is later)")
    print("  Preview response:   ID 123456789, Status 'Survey Preview' (should be ignored)")
    zeros = [s for s, _ in students if s.startswith("0")]
    if zeros:
        print(f"  Leading zero lost:  {', '.join(zeros)} (should still match)")


if __name__ == "__main__":
    main()

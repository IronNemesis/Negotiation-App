"""Salary negotiation outcomes sheet: two Qualtrics exports from before class plus the in-class Google Form.

Called by brs.py for --simulation SL (prepare, build, missing). Reuses brs's price cleaning and response matching;
adds what the Salary sheet needs:
  - every money value is a total salary: a number below `raise_below` is read as a raise and added to `base_salary`
  - the Google Form (one row per match, typed by students) is matched to the roster's matches
  - her layout: Supervisor on the left, BASE / RAISE / TOTAL in the middle, Subordinate mirrored on the right
"""
from __future__ import annotations

import csv
import re
import shutil
import tempfile
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import brs
import nps_shared as ns
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font

MONEY_FORMAT = '"$"#,##0'
PLAN = ["INITIAL", "TARGET", "RESERVATION"]


# ---------------------------------------------------------------- money

@dataclass
class Money:
    value: object            # a number, the text as typed, or None
    flag: str = ""           # "", "raise" (read as a raise), "corrected", "unclear", "out_of_range", "mismatch"
    typed: str = ""
    notes: list[str] = field(default_factory=list)


def money(text: str, kind: str, sim: dict) -> Money:
    """kind: "salary" (a raise or a total; shown as a total), "total" (always a total), "amount" (a raise as typed)."""
    typed = ns.clean(text)
    if not typed:
        return Money(None)
    if re.fullmatch(r"no\s*deal", typed, re.I):
        return Money(typed, "", typed)
    a = brs.clean_price(typed, 0, 10 ** 9)
    if not isinstance(a.value, (int, float)):
        return Money(a.value, "unclear", typed, [f"Kept exactly as typed: {typed}"])
    value, flag, notes = a.value, a.flag, []
    if flag == "corrected":
        notes.append(f"Typed: {typed}\nRead as {value:,}.")
    base, cutoff = sim["base_salary"], sim["raise_below"]
    if kind == "salary" and value < cutoff:
        notes.append(f"Typed as a raise of ${value:,}; shown as the total salary (${base:,} base + raise).")
        value = base + value
        flag = flag or "raise"
    low, high = sim["salary_range"]
    if kind in ("salary", "total") and not low <= value <= high:
        flag = "out_of_range"
        notes.append(f"Typed: {typed}. Outside the usual salary range; check it.")
    return Money(value, flag, typed, notes)


def minutes(text: str) -> Money:
    typed = ns.clean(text)
    if not typed or typed.isdigit():
        return Money(int(typed) if typed else None, "", typed)
    m = re.fullmatch(r"(\d+)\s*(min|mins|minutes)\.?", typed, re.I)
    if m:
        return Money(int(m.group(1)), "corrected", typed, [f"Typed: {typed}\nRead as {m.group(1)}."])
    return Money(typed, "unclear", typed, [f"Kept exactly as typed: {typed}"])


# ---------------------------------------------------------------- reading

def read_csv(path: Path) -> list[list[str]]:
    if not path.exists():
        raise ns.SkillError("file_missing", f"Could not find {path.name} in {path.parent}.")
    try:
        with path.open(newline="", encoding="utf-8-sig") as f:
            return list(csv.reader(f))
    except PermissionError:
        raise ns.SkillError("file_locked", f"{path.name} is open in Excel. Close it and try again.")


def read_surveys(paths: list[str], sim: dict, roster: Path, code: str, name_column: str | None):
    """Both Qualtrics exports, told apart by their questions. Returns {role: (path, raw rows, responses, removed)}."""
    out = {}
    for text in paths:
        path = Path(text).expanduser().resolve()
        rows = read_csv(path)
        questions = " ".join(rows[1] if len(rows) > 1 else []).lower()
        roles = [role for role, cfg in sim["surveys"].items() if cfg["detect"].lower() in questions]
        if len(roles) != 1:
            raise ns.SkillError("survey_layout", f"{path.name} doesn't look like the Supervisor or the Subordinate "
                                                 "Salary survey. The survey may have been edited.")
        role = roles[0]
        if role in out:
            raise ns.SkillError("bad_option", f"Two {role.title()} exports were given: {out[role][0].name} and "
                                              f"{path.name}. Pass one Supervisor and one Subordinate export.")
        cfg = dict(sim["surveys"][role], role_code=str(sim["role_codes"][role]))
        raw, _, responses, removed, _ = brs.read_export(path, {"survey": cfg, "price_range": [0, 10 ** 9]}, roster,
                                                        code, name_column)
        for r in responses:   # read again as salaries: brs read them as plain prices
            r.prices = {q: money(next(row for row in raw if len(row) > 8 and row[8] == r.rid)[
                raw[0].index(cfg["prices"][q])], "salary", sim) for q in PLAN}
        out[role] = (path, raw, responses, removed)
    return out


@dataclass
class FormRow:
    row: int                 # row number in the CSV as she sees it in Excel (the header is row 1)
    stamp: datetime | None
    values: dict[str, str]
    match: int | None = None
    note: str = ""


def read_form(path: Path, sim: dict) -> tuple[list[list[str]], list[FormRow]]:
    rows = read_csv(path)
    if not rows:
        raise ns.SkillError("survey_empty", f"{path.name} is empty.")
    header = [ns.clean(h).lower() for h in rows[0]]
    cols = {}
    for key, prefix in sim["form"].items():
        hits = [i for i, h in enumerate(header) if h.startswith(prefix.lower())]
        if not hits:
            raise ns.SkillError("survey_layout", f"The Form export has no column starting with '{prefix}'. The Form "
                                                 "may have been edited; the skill's settings need updating.",
                                {"columns": rows[0]})
        cols[key] = hits[0]
    out = []
    for n, row in enumerate(rows[1:], start=2):
        row = row + [""] * (len(header) - len(row))
        if not any(ns.clean(v) for v in row):
            continue
        values = {key: row[i] for key, i in cols.items()}
        out.append(FormRow(n, parse_stamp(values["timestamp"]), values))
    return rows, out


def parse_stamp(text: str) -> datetime | None:
    for fmt in ("%m/%d/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M"):
        try:
            return datetime.strptime(ns.clean(text), fmt)
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------- matching the Form to matches

def typed_people(first: str, last: str) -> list[tuple[str, str]]:
    """'Lena and Siobhan' / 'Ortiz and O'Neill' -> two people."""
    firsts = [x.strip() for x in re.split(r"\s+and\s+|\s*&\s*|\s*,\s*", ns.clean(first)) if x.strip()]
    lasts = [x.strip() for x in re.split(r"\s+and\s+|\s*&\s*|\s*,\s*", ns.clean(last)) if x.strip()]
    if len(firsts) == len(lasts):
        return list(zip(firsts, lasts))
    return [(ns.clean(first), ns.clean(last))]


def match_form(form: list[FormRow], groups: dict, args) -> tuple[dict[int, FormRow], list[str], list[dict]]:
    """Each Form row goes to one match: by the names typed when they agree on one match, otherwise by the match
    number. Returns {match: the row used}, notes, unresolved rows."""
    people = {role: [p for g in groups.values() for p in g if p.role == role] for role in ("SUPERVISOR", "SUBORDINATE")}
    pools = {role: brs.pool_students(ps) for role, ps in people.items()}
    match_of = {p.name: p.match for g in groups.values() for p in g}
    placed = {}
    for spec in args.form_match:
        row, _, match = spec.partition("=")
        if not row.strip().isdigit() or not match.strip().isdigit() or int(match) not in groups:
            raise ns.SkillError("bad_decision", f"--form-match {spec}: use the Form row number and a match number "
                                                "still in use, e.g. 4=12.")
        placed[int(row)] = int(match)
    skip = {int(x) for x in args.form_skip if x.strip().isdigit()}
    notes, unresolved, by_match = [], [], {}
    for fr in form:
        if fr.row in skip:
            notes.append(f"Form row {fr.row} left out by the instructor.")
            continue
        typed_match = ns.clean(fr.values["match"])
        number = re.match(r"\d+", typed_match)
        number = int(number.group()) if number else None
        named = set()
        for role, f_key, l_key in (("SUPERVISOR", "sup_first", "sup_last"), ("SUBORDINATE", "sub_first", "sub_last")):
            for first, last in typed_people(fr.values[f_key], fr.values[l_key]):
                res = ns.resolve_student(first, last, pools[role])
                if res.confident:
                    named.add(match_of[res.student.name])
        if fr.row in placed:
            fr.match = placed[fr.row]
            notes.append(f"Form row {fr.row} put on match {fr.match} by the instructor.")
        elif len(named) == 1 and (number is None or number not in groups or next(iter(named)) != number):
            fr.match = next(iter(named))
            notes.append(f"Form row {fr.row}: the match number was typed as '{typed_match}'; the names are in "
                         f"match {fr.match}, so it's used there.")
        elif number in groups and (not named or named == {number}):
            fr.match = number
            if typed_match != str(number):
                notes.append(f"Form row {fr.row}: the match number was typed as '{typed_match}'; read as {number}.")
        else:
            unresolved.append({"type": "form_row", "form_row": fr.row, "typed_match": typed_match,
                               "typed_names": f"{fr.values['sup_first']} {fr.values['sup_last']} / "
                                              f"{fr.values['sub_first']} {fr.values['sub_last']}".strip(),
                               "reason": "the match number and the names don't point to one match" if named
                               else "the match number isn't a match in use and the names don't match anyone",
                               "suggestions": sorted(named)})
            continue
        by_match.setdefault(fr.match, []).append(fr)
    used = {}
    for match, frs in by_match.items():
        frs.sort(key=lambda f: (f.stamp or datetime.min, f.row))
        used[match] = frs[-1]
        if len(frs) > 1:
            notes.append(f"Match {match} has {len(frs)} Form rows (rows {', '.join(str(f.row) for f in frs)}); "
                         f"the latest (row {frs[-1].row}) is used.")
    return used, notes, unresolved


# ---------------------------------------------------------------- the sheet

COLUMNS = [  # (header, source, key)  source: "form" yellow / "plan" white / "name"
    ("MATCH", "form", "match"),
    ("SUPERVISOR FIRST NAME", "form", "first"), ("SUPERVISOR LAST NAME", "form", "last"),
    ("ACTUAL INITIAL OFFER", "form", "sup_initial"),
    ("INITIAL OFFER", "plan", "INITIAL"), ("TARGET", "plan", "TARGET"), ("RESERVATION", "plan", "RESERVATION"),
    ("BASE", "form", "base"), ("RAISE", "form", "raise"), ("TOTAL", "form", "total"),
    ("RESERVATION", "plan", "RESERVATION"), ("TARGET", "plan", "TARGET"), ("INITIAL OFFER", "plan", "INITIAL"),
    ("ACTUAL INITIAL OFFER", "form", "sub_initial"),
    ("SUBORDINATE FIRST NAME", "form", "first"), ("SUBORDINATE LAST NAME", "form", "last"),
    ("Did the SUBORDINATE make a credible threat to walk away?", "form", "threat"),
    ("How many minutes did you negotiate for?", "form", "minutes"),
    ("To what extent was THE SUPERVISOR'S approach collaborative or competitive? (The subordinate must answer this "
     "question)", "form", "sup_approach"),
    ("To what extent was THE SUBORDINATE'S approach collaborative or competitive? (The supervisor must answer this "
     "question)", "form", "sub_approach"),
    ("Any other terms of the agreement:", "form", "terms"),
]
SUP_COLS = {"first": 2, "last": 3, "INITIAL": 5, "TARGET": 6, "RESERVATION": 7}
SUB_COLS = {"RESERVATION": 11, "TARGET": 12, "INITIAL": 13, "first": 15, "last": 16}
MATCH_COLS = {"sup_initial": 4, "base": 8, "raise": 9, "total": 10, "sub_initial": 14, "threat": 17, "minutes": 18,
              "sup_approach": 19, "sub_approach": 20, "terms": 21}
MONEY_COLS = {4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14}


def form_values(fr: FormRow, sim: dict) -> dict[str, Money]:
    v = fr.values
    out = {"sup_initial": money(v["sup_initial"], "salary", sim), "sub_initial": money(v["sub_initial"], "salary", sim),
           "base": money(v["base"], "total", sim), "raise": money(v["raise"], "amount", sim),
           "total": money(v["total"], "total", sim), "minutes": minutes(v["minutes"])}
    for key in ("threat", "sup_approach", "sub_approach", "terms"):
        out[key] = Money(ns.clean(v[key]) or None, "", ns.clean(v[key]))
    base, raise_, total = out["base"].value, out["raise"].value, out["total"].value
    if all(isinstance(x, (int, float)) for x in (base, raise_, total)) and base + raise_ != total:
        out["total"].flag = "mismatch"
        out["total"].notes.append(f"Base + raise = ${base + raise_:,}, but the total was typed as {out['total'].typed}.")
    return out


def put(ws, expected, row, col, m: Money | object):
    cell = ws.cell(row, col)
    if not isinstance(m, Money):
        cell.value = m
        expected[(row, col)] = m
        return
    cell.value = m.value
    expected[(row, col)] = m.value
    if col in MONEY_COLS and isinstance(m.value, (int, float)):
        cell.number_format = MONEY_FORMAT
    if m.flag == "corrected":
        cell.fill = brs.fill(brs.ORANGE)
    elif m.flag in ("unclear", "out_of_range", "mismatch"):
        cell.fill = brs.fill(brs.LIGHT_RED)
    if m.notes:
        cell.comment = Comment("\n".join(m.notes), brs.AUTHOR)


def write_sheet(wb, title: str, groups, absent, used: dict, sim: dict) -> dict:
    ws = wb.active
    ws.title = title
    for c, (header, source, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, c, header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if source == "form":
            cell.fill = brs.fill(brs.YELLOW)
    expected = {}

    def person(row, p, cols):
        put(ws, expected, row, cols["first"], p.first)
        put(ws, expected, row, cols["last"], p.last)
        if p.response:
            for q in PLAN:
                put(ws, expected, row, cols[q], p.response.prices[q])
        else:
            ws.cell(row, cols["first"]).comment = Comment("No survey response.", brs.AUTHOR)
        if p.role_conflict:
            ws.cell(row, cols["first"]).fill = brs.fill(brs.LIGHT_RED)
            ws.cell(row, cols["first"]).comment = Comment(p.role_conflict.replace("Chose", "Took the") +
                                                          " (the other role's survey).", brs.AUTHOR)

    row = 2
    for index, (match, people) in enumerate(sorted(groups.items())):
        sups = [p for p in people if p.role == "SUPERVISOR"]
        subs = [p for p in people if p.role == "SUBORDINATE"]
        height = max(len(sups), len(subs))
        for k in range(height):
            if index % 2 == 0:
                for col in range(1, len(COLUMNS) + 1):
                    ws.cell(row, col).fill = brs.fill(brs.BAND)
            if k == 0:
                put(ws, expected, row, 1, match)
                if match in used:
                    for key, value in form_values(used[match], sim).items():
                        put(ws, expected, row, MATCH_COLS[key], value)
                else:
                    ws.cell(row, 1).comment = Comment("No Form response for this match.", brs.AUTHOR)
            if k < len(sups):
                person(row, sups[k], SUP_COLS)
            if k < len(subs):
                person(row, subs[k], SUB_COLS)
            row += 1
    if absent:   # kept for later research, as on the Used Car sheet
        row += 1
        for p in absent:
            put(ws, expected, row, 1, brs.ABSENT_LABEL)
            person(row, p, SUB_COLS if p.role == "SUBORDINATE" else SUP_COLS)
            row += 1
    legend = row + 1
    for offset, (color, text) in enumerate([
            (brs.YELLOW, "From the in-class Google Form (other columns: the Qualtrics survey before class)"),
            (brs.ORANGE, "Typo corrected by Claude; the original answer is in the cell comment"),
            (brs.LIGHT_RED, "Unclear or unusual answer kept as typed, or a total that doesn't equal base + raise")]):
        ws.cell(legend + offset, 1).fill = brs.fill(color)
        ws.cell(legend + offset, 2, text)
    ws.cell(legend + 3, 2, f"Money is shown as total salary: a raise is added to the ${sim['base_salary']:,} base. "
                           "Hover over a value to see what was typed.")
    widths = [7, 12, 14] + [11] * 11 + [12, 14, 14, 11, 18, 18, 30]
    for c, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = width
    ws.row_dimensions[1].height = 75
    ws.freeze_panes = "B2"
    return expected


def write_raw(wb, title: str, rows):
    ws = wb.create_sheet(title)
    for r in rows:
        ws.append(r)


# ---------------------------------------------------------------- command

def run(args, mode: str, sim: dict) -> None:
    started = time.time()
    brs.SETTINGS_SIM.clear()   # brs.py runs as __main__, so this imported copy needs the settings too
    brs.SETTINGS_SIM.update(sim)
    roster = Path(args.roster).expanduser().resolve()
    surveys = read_surveys(args.csv, sim, roster, args.simulation, args.name_column)
    students, stage, groups, absent, warnings = brs.load_matches(roster, args.simulation, sim)
    responses = [r for _, _, rs, _ in surveys.values() for r in rs]
    removed = [x for _, _, _, rm in surveys.values() for x in rm]
    decisions, unresolved, left_off = brs.match_responses(responses, groups, absent, students, args)
    summary = brs.summarize(stage, groups, absent, removed, responses, decisions, unresolved, left_off, warnings)
    summary["cleaning"] = {"corrected": [], "unclear": [], "out_of_range": []}
    raises = 0
    owner = {p.response.rid: p.name for g in list(groups.values()) + [absent] for p in g if p.response}
    for r in responses:
        if r.rid not in owner:   # a duplicate that wasn't used, or a response not matched yet
            continue
        for q, m in r.prices.items():
            raises += m.flag == "raise" or (m.flag == "corrected" and any("raise" in n for n in m.notes))
            if m.flag in summary["cleaning"]:
                summary["cleaning"][m.flag].append({"student": owner[r.rid],
                                                    "question": q, "typed": m.typed, "now": m.value})
    summary["raises_shown_as_total"] = raises
    summary["assumptions"] = {"base_salary": sim["base_salary"],
                              "a_number_below_this_is_a_raise": sim["raise_below"]}
    missing_roles = [role for role in sim["roles"] if role not in surveys]
    if missing_roles:
        summary["missing_exports"] = [f"{role.title()} survey" for role in missing_roles]
    result = {"ok": True, **summary}

    form_rows, used = None, {}
    if args.form:
        form_path = Path(args.form).expanduser().resolve()
        form_rows, form = read_form(form_path, sim)
        used, form_notes, form_unresolved = match_form(form, groups, args)
        result["form"] = {"rows": len(form), "matches_with_a_row": len(used), "notes": form_notes,
                          "matches_without_a_row": sorted(set(groups) - set(used))}
        problems = {"mismatch": [], "unclear": [], "corrected": [], "out_of_range": []}
        for match, fr in sorted(used.items()):
            for key, m in form_values(fr, sim).items():
                if m.flag in problems:
                    problems[m.flag].append({"match": match, "column": key, "typed": m.typed, "now": m.value})
        result["form"]["cleaning"] = problems
        result["unresolved"] = result["unresolved"] + form_unresolved
        result["ready_to_build"] = not result["unresolved"]
    elif mode == "build":
        result["form"] = "No Form export given: the yellow (in-class) columns are left blank."

    if mode == "missing":
        listed = brs.missing_list(roster, args.simulation, groups, unresolved)
        ns.write_log(roster.parent, "T10 Survey non-respondents (SL)",
                     f"{len(listed['missing'])} of {result['students']} students without a response.")
        ns.emit({"ok": True, "matches_from": result["matches_from"], "students": result["students"], **listed})
        return
    if mode != "build":
        ns.write_log(roster.parent, "T6-T7 Prepare salary sheet (SL)",
                     f"{result['students']} students, {len(responses)} responses kept, "
                     f"{len(result['unresolved'])} unresolved, matches from the {stage} table.")
        ns.emit(result)
        return
    if missing_roles:
        raise ns.SkillError("export_missing", f"The {' and '.join(summary['missing_exports'])} export is needed to "
                                              "build the sheet.")
    if result["unresolved"]:
        raise ns.SkillError("unresolved_responses", "Some survey responses or Form rows need the instructor's "
                                                    "decision first.", {"unresolved": result["unresolved"]})
    course = args.course or brs.course_from_roster(roster)
    if not course:
        raise ns.SkillError("course_unknown", "Could not tell the course and term from the roster file name. "
                                              "Pass --course, e.g. 'BUS 4489 F26'.")
    folder = (Path(args.form).expanduser().resolve() if args.form else surveys[sim["roles"][0]][0]).parent
    out = folder / f"{course} Week {args.week} - {sim['file_title']}.xlsx"
    previous = ns.backup(out, "replaced") if out.exists() else None

    wb = Workbook()
    expected = write_sheet(wb, brs.safe_title(f"Week {args.week} {sim['name']}"), groups, absent, used, sim)
    brs.write_roles_tab(wb, sim["roles"], groups)
    for role in sim["roles"]:
        write_raw(wb, f"{role.title()} survey", surveys[role][1])
    if form_rows is not None:
        write_raw(wb, "Form", form_rows)
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / out.name
        wb.save(tmp_path)
        check = load_workbook(tmp_path).worksheets[0]
        mismatched = [check.cell(r, c).coordinate for (r, c), v in expected.items() if check.cell(r, c).value != v]
        if mismatched:
            raise ns.SkillError("readback_mismatch", "Some values did not save correctly; the sheet was not created.",
                                {"cells": mismatched})
        try:
            shutil.copyfile(tmp_path, out)
        except PermissionError:
            raise ns.SkillError("file_locked", f"{out.name} is open in Excel. Close it and try again.")
    result.update({"workbook": str(out), "previous_version_backup": str(previous) if previous else None,
                   "seconds": round(time.time() - started, 1)})
    warning = ns.write_log(roster.parent, "T8 Build salary sheet (SL)",
                           f"{out.name} from the {stage} table: {result['students']} students, "
                           f"{len(result['no_response'])} without a survey response, "
                           f"{len(used)} matches with a Form row. Previous version: "
                           f"{previous.name if previous else 'none'}.")
    if warning:
        result["log_warning"] = warning
    ns.emit(result)

# GENERATED COPY: do not edit here. Edit negotiation_agent/shared/nps_shared.py
# and run negotiation_agent/shared/build_skills.py to update every skill.
"""Shared code for the negotiation-pairing and bargaining-range-sheet skills.

Edit this file here only. `build_skills.py` copies it into each skill's
scripts/ folder, so every skill folder works on its own when uploaded.

It reads the instructor's roster workbook following
negotiation_agent/docs/workbook-conventions.md: the class list on the
attendance tab, the final and draft tables on a simulation tab, and the pair
IDs on every simulation tab. It also holds name matching, backups, the log,
and the JSON output every command uses.
"""
from __future__ import annotations

import json
import re
import shutil
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    print(json.dumps({"ok": False, "error": "missing_dependency",
                      "message": "The Python package openpyxl is not installed. Install it with: pip install openpyxl"}))
    sys.exit(1)

LOG_NAME = "Negotiation Log.xlsx"
BACKUP_FOLDER = "Backups"
PAIR_ID = re.compile(r"^([A-Z]{2})_(\d+)$")
CLASS_NAME = re.compile(r"^[^,@]+,\s*[^,@]+$")  # "Last, First"


class SkillError(Exception):
    """A failure the instructor needs to hear about in plain language."""

    def __init__(self, code: str, message: str, details: dict | None = None):
        super().__init__(message)
        self.code, self.message, self.details = code, message, details or {}


# ---------------------------------------------------------------- output

def emit(result: dict) -> None:
    print(json.dumps(result, indent=2, ensure_ascii=False, default=str))


def run_cli(main) -> None:
    """Run a command's main(); turn SkillError into a JSON failure with exit code 2."""
    try:
        main()
    except SkillError as exc:
        emit({"ok": False, "error": exc.code, "message": exc.message, **exc.details})
        sys.exit(2)


# ---------------------------------------------------------------- text and names

def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize(text) -> str:
    """Compare names the way a person would: ignore accents, case, and extra spaces."""
    text = unicodedata.normalize("NFKD", clean(text))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"['’‘`]", "", text.casefold())  # O'Neill, O’Neill (phone keyboards), and ONeill all match
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9, -]", "", text)).strip()


def split_class_name(name: str) -> tuple[str, str]:
    """'Wendel Jeansson, Rasmus' -> ('Rasmus', 'Wendel Jeansson')."""
    last, _, first = clean(name).partition(",")
    return first.strip(), last.strip()


def name_key(first: str, last: str) -> str:
    return f"{normalize(last)}, {normalize(first)}"


def class_name_key(name: str) -> str:
    first, last = split_class_name(name)
    return name_key(first, last)


def looks_like_class_name(value) -> bool:
    return isinstance(value, str) and bool(CLASS_NAME.match(value.strip()))


# Common English nicknames. Matching also accepts one first name being the start of the other ("Dan" / "Daniel").
NICKNAME_GROUPS = [
    {"andrew", "andy", "drew"}, {"alexander", "alex", "xander"}, {"alexandra", "alex", "lexi", "sasha"},
    {"anthony", "tony"}, {"benjamin", "ben", "benji"}, {"catherine", "katherine", "kathryn", "kate", "katie", "kat", "cathy"},
    {"christopher", "chris"}, {"christina", "christine", "chris", "tina"}, {"daniel", "dan", "danny"},
    {"danielle", "dani"}, {"david", "dave"}, {"edward", "ed", "eddie", "ted"}, {"elizabeth", "liz", "beth", "eliza", "lizzie"},
    {"emily", "em"}, {"gabriel", "gabe"}, {"gabriela", "gabriella", "gabby"}, {"isabella", "isabel", "bella", "izzy"},
    {"jacob", "jake"}, {"james", "jim", "jimmy", "jamie"}, {"jennifer", "jen", "jenny"}, {"jessica", "jess"},
    {"jonathan", "jon", "john", "johnny"}, {"joseph", "joe", "joey"}, {"joshua", "josh"}, {"katelyn", "kate", "katie"},
    {"lawrence", "larry"}, {"madison", "maddie", "maddy"}, {"margaret", "maggie", "meg", "peggy"},
    {"matthew", "matt"}, {"michael", "mike", "mikey"}, {"michelle", "shelly"}, {"nathaniel", "nathan", "nate"},
    {"nicholas", "nick", "nicky"}, {"nicole", "nikki"}, {"patrick", "pat"}, {"rebecca", "becca", "becky"},
    {"richard", "rich", "rick", "dick"}, {"robert", "rob", "bob", "bobby", "robbie"}, {"samantha", "sam", "sammie"},
    {"samuel", "sam", "sammy"}, {"stephanie", "steph"}, {"steven", "stephen", "steve"}, {"theodore", "theo", "ted"},
    {"thomas", "tom", "tommy"}, {"victoria", "tori", "vicky"}, {"william", "will", "bill", "billy", "liam"},
    {"zachary", "zach", "zack"},
]


def first_names_compatible(a: str, b: str) -> bool:
    a, b = normalize(a), normalize(b)
    if not a or not b:
        return False
    if a == b:
        return True
    if len(min(a, b, key=len)) >= 3 and (a.startswith(b) or b.startswith(a)):
        return True
    return any(a in group and b in group for group in NICKNAME_GROUPS)


# ---------------------------------------------------------------- workbook access

def open_workbook(path: Path, data_only: bool = False):
    path = Path(path)
    if not path.exists():
        raise SkillError("file_missing", f"Could not find {path.name} in {path.parent}.", {"path": str(path)})
    try:
        return load_workbook(path, data_only=data_only)
    except PermissionError:
        raise SkillError("file_locked", f"{path.name} is open in Excel or being synced by OneDrive. "
                                        "Close it and try again.", {"path": str(path)})
    except Exception as exc:
        raise SkillError("file_unreadable", f"Could not read {path.name}: {exc}", {"path": str(path)})


def save_workbook(wb, path: Path) -> None:
    try:
        wb.save(path)
    except PermissionError:
        raise SkillError("file_locked", f"{Path(path).name} is open in Excel or being synced by OneDrive. "
                                        "Close it and try again.", {"path": str(path)})


def set_cell(ws, row: int, column: int, value) -> None:
    """Write a cell, including clearing it. Use this instead of ws.cell(row, column, value):
    openpyxl ignores value=None in that call, so a cell meant to be cleared would keep its old value."""
    ws.cell(row, column).value = value


def backup(path: Path, label: str) -> Path:
    """Copy a file into the Backups folder next to it before it is changed."""
    path = Path(path)
    folder = path.parent / BACKUP_FOLDER
    folder.mkdir(exist_ok=True)
    target = folder / f"{path.stem} {datetime.now():%Y-%m-%d %H%M%S} {label}{path.suffix}"
    shutil.copy2(path, target)
    return target


def write_log(folder: Path, step: str, details: str) -> str | None:
    """Append one plain-language line to the log next to the instructor's files.

    Returns a warning instead of failing, so a log that is open in Excel never blocks the real work.
    """
    path = Path(folder) / LOG_NAME
    try:
        if path.exists():
            wb = load_workbook(path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Log"
            ws.append(["Time", "Step", "Details"])
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="D9E7DF")
            ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 20, 34
            ws.column_dimensions["C"].width = 110
        ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), step, details])
        wb.save(path)
        return None
    except PermissionError:
        return f"{LOG_NAME} is open in Excel, so this entry was not logged: {step}: {details}"


# ---------------------------------------------------------------- class list

@dataclass
class Student:
    name: str          # exactly as on the class list, "Last, First"
    first: str
    last: str
    email: str
    row: int

    @property
    def key(self) -> str:
        return name_key(self.first, self.last)


def find_header_row(ws, required: list[str], max_rows: int = 10) -> tuple[int, dict[str, int]] | None:
    """Return (row number, {normalized header: column number}) for the first row holding every required header."""
    for r in range(1, min(ws.max_row, max_rows) + 1):
        headers = {normalize(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1)
                   if clean(ws.cell(r, c).value)}
        if all(any(req in h for h in headers) for req in required):
            return r, headers
    return None


def read_class_list(wb, tab: str | None = None) -> tuple[str, list[Student], list[str]]:
    """Read the class list from the attendance tab. Returns (tab name, students, warnings)."""
    candidates = [tab] if tab else [name for name in wb.sheetnames if name.lower().startswith("att")]
    if tab and tab not in wb.sheetnames:
        raise SkillError("class_tab_missing", f"There is no tab named '{tab}'.", {"tabs": wb.sheetnames})
    found = []
    for name in candidates:
        ws = wb[name]
        header = find_header_row(ws, ["student name", "email"])
        if not header:
            continue
        hr, cols = header
        c_name = next(c for h, c in cols.items() if "student name" in h)
        c_email = next(c for h, c in cols.items() if "email" in h)
        students, warnings = [], []
        for r in range(hr + 1, ws.max_row + 1):
            raw = clean(ws.cell(r, c_name).value)
            if not raw:
                continue
            if "," not in raw:
                warnings.append(f"Row {r} on '{name}': '{raw}' is not in 'Last, First' form and was skipped.")
                continue
            first, last = split_class_name(raw)
            email = clean(ws.cell(r, c_email).value).lower()
            if not email:
                warnings.append(f"{raw} has no email on '{name}'.")
            students.append(Student(raw, first, last, email, r))
        if students:
            found.append((name, students, warnings))
    if not found:
        raise SkillError("class_list_missing", "Could not find a class list: no attendance tab has "
                                               "'Student Name' and 'Email' headers with names under them.",
                         {"tabs": wb.sheetnames})
    if len(found) > 1:
        raise SkillError("class_tab_ambiguous", "More than one attendance tab has students on it. "
                                                "Say which one is the current class list.",
                         {"tabs": [f[0] for f in found]})
    name, students, warnings = found[0]
    keys = {}
    for s in students:
        if s.key in keys:
            raise SkillError("duplicate_student", f"'{s.name}' appears twice on '{name}' "
                                                  f"(rows {keys[s.key]} and {s.row}).")
        keys[s.key] = s.row
    return name, students, warnings


@dataclass
class Resolution:
    student: Student | None
    how: str              # "exact", "email", "nickname", "likely" (needs confirmation), or "none"
    candidates: list[Student] = field(default_factory=list)

    @property
    def confident(self) -> bool:
        return self.student is not None and self.how in ("exact", "email", "nickname")


def resolve_student(first: str, last: str, students: list[Student], email: str = "") -> Resolution:
    """Link a name (and optional email) to one class-list student, or say why it can't.

    Her own tabs and student-typed surveys use variants: "Zach" for "Zachary", a preferred first name, a
    personal email. Only exact names, class-list emails, and known nicknames with a unique last name count as
    confident. A unique last name with a different first name is "likely" and must be confirmed.
    """
    key = name_key(first, last)
    exact = [s for s in students if s.key == key]
    if len(exact) == 1:
        return Resolution(exact[0], "exact")
    if email:
        by_email = [s for s in students if s.email and s.email == clean(email).lower()]
        if len(by_email) == 1:
            return Resolution(by_email[0], "email")
    same_last = [s for s in students if normalize(s.last) == normalize(last)]
    compatible = [s for s in same_last if first_names_compatible(s.first, first)]
    if len(same_last) == 1 and len(compatible) == 1:
        return Resolution(compatible[0], "nickname")
    if len(same_last) == 1:
        return Resolution(same_last[0], "likely", same_last)
    candidates = compatible or same_last or [s for s in students if first_names_compatible(s.first, first)]
    return Resolution(None, "none", candidates[:3])


def resolve_class_name(name: str, students: list[Student], email: str = "") -> Resolution:
    first, last = split_class_name(name)
    return resolve_student(first, last, students, email)


# ---------------------------------------------------------------- simulation tabs

@dataclass
class Member:
    match: int | None   # None for a student on an extra row (they belong to the match above)
    role: str
    name: str
    email: str = ""
    pair_id: str = ""
    row: int = 0


@dataclass
class RoleColumns:
    role: str
    name: int
    email: int | None = None
    prior_ids: dict[str, int] = field(default_factory=dict)  # earlier simulation code -> column
    current_id: int | None = None


@dataclass
class SimTab:
    title: str
    code: str
    roles: list[str]
    final_header: int
    final_columns: list[RoleColumns]
    match_column: int
    draft_header: int
    draft_columns: dict[str, int]       # role -> column
    draft_match_column: int
    final_members: list[Member]
    absent: list[Member]
    final_prefilled: list[int]          # rows that hold only a prefilled match number
    draft_members: list[Member]
    draft_prefilled: list[int]

    def groups(self, members: list[Member]) -> dict[int, list[Member]]:
        """Group members by match number; extra rows join the match above them."""
        out: dict[int, list[Member]] = {}
        current = None
        for m in members:
            if m.match is not None:
                current = m.match
            if current is None:
                raise SkillError("orphan_row", f"{m.name} (row {m.row} on '{self.title}') is on an extra row "
                                               "with no match above it.")
            out.setdefault(current, []).append(m)
        return out


def _role_header_rows(ws, roles: list[str]) -> list[tuple[int, dict[str, int]]]:
    wanted = [r.upper() for r in roles]
    rows = []
    for r in range(1, ws.max_row + 1):
        cells = {clean(ws.cell(r, c).value).upper(): c for c in range(1, ws.max_column + 1)
                 if clean(ws.cell(r, c).value)}
        if all(role in cells for role in wanted):
            rows.append((r, cells))
    return rows


def read_sim_tab(wb, code: str, roles: list[str]) -> SimTab:
    """Locate and read the final (top) and draft (bottom) tables of a simulation tab."""
    if code not in wb.sheetnames:
        raise SkillError("sim_tab_missing", f"There is no '{code}' tab in the workbook. Prepare the tab first.",
                         {"tabs": wb.sheetnames})
    ws = wb[code]
    header_rows = _role_header_rows(ws, roles)
    final = [(r, cells) for r, cells in header_rows if "EMAIL" in cells]
    draft = [(r, cells) for r, cells in header_rows if "EMAIL" not in cells and "MATCH" in cells]
    if len(final) != 1 or len(draft) != 1:
        raise SkillError("sim_tab_layout",
                         f"The '{code}' tab should have one header row with {' and '.join(roles)} and Email "
                         f"(the final table) and one with MATCH, {' and '.join(roles)} (the draft table). "
                         f"Found {len(final)} and {len(draft)}.")
    (fr, fcells), (dr, dcells) = final[0], draft[0]
    if dr <= fr:
        raise SkillError("sim_tab_layout", f"On the '{code}' tab the draft table (MATCH header) should be below "
                                           "the final table.")

    # Final table columns: role name, its Email, then earlier codes, then this simulation's pair ID.
    role_starts = sorted((fcells[r.upper()], r) for r in roles)
    final_columns = []
    for i, (col, role) in enumerate(role_starts):
        limit = role_starts[i + 1][0] if i + 1 < len(role_starts) else ws.max_column + 1
        rc = RoleColumns(role=role, name=col)
        c = col + 1
        header_says_email = "email" in normalize(ws.cell(fr, c).value)
        # Her headers are sometimes blank above an email column (the SL tab's second role), so look at the cells too.
        holds_emails = not clean(ws.cell(fr, c).value) and any(
            "@" in clean(ws.cell(r, c).value) for r in range(fr + 1, min(fr + 40, ws.max_row + 1)))
        if c < limit and (header_says_email or holds_emails):
            rc.email = c
            c += 1
        while c < limit:
            header = clean(ws.cell(fr, c).value).upper()
            if header == code or (not header and rc.current_id is None):
                rc.current_id = c
                c += 1
                break
            if re.fullmatch(r"[A-Z]{2}", header):
                rc.prior_ids[header] = c
                c += 1
                continue
            break
        if rc.email is None or rc.current_id is None:
            raise SkillError("sim_tab_layout", f"On the '{code}' tab, the {role} column in the final table should be "
                                               f"followed by an Email column and a pair-ID column.")
        final_columns.append(rc)
    match_col = fcells.get("MATCH", 1)
    draft_cols = {r: dcells[r.upper()] for r in roles}
    draft_match_col = dcells["MATCH"]

    def is_number(v) -> bool:
        return isinstance(v, (int, float)) or (isinstance(v, str) and v.strip().isdigit())

    # Final table rows, between its header and the draft header.
    final_members, absent, final_prefilled = [], [], []
    seen_names, gap_after_names = False, False
    for r in range(fr + 1, dr):
        match_val = ws.cell(r, match_col).value
        row_members = []
        for rc in final_columns:
            name = clean(ws.cell(r, rc.name).value)
            if name:
                row_members.append(Member(None, rc.role, name, clean(ws.cell(r, rc.email).value).lower(),
                                          clean(ws.cell(r, rc.current_id).value), r))
        if not row_members:
            if is_number(match_val):
                final_prefilled.append(r)
            elif seen_names:
                gap_after_names = True
            continue
        if gap_after_names or any(m.pair_id.endswith("_0") for m in row_members):
            absent.extend(row_members)
            continue
        seen_names = True
        match = int(match_val) if is_number(match_val) else None
        for m in row_members:
            m.match = match
        final_members.extend(row_members)

    # Draft table rows, until the first fully blank row after names start.
    draft_members, draft_prefilled = [], []
    started = False
    for r in range(dr + 1, ws.max_row + 1):
        match_val = ws.cell(r, draft_match_col).value
        names = [(role, clean(ws.cell(r, col).value)) for role, col in draft_cols.items()
                 if clean(ws.cell(r, col).value)]
        if not names:
            if is_number(match_val):
                draft_prefilled.append(r)
                continue
            if started:
                break
            continue
        started = True
        match = int(match_val) if is_number(match_val) else None
        draft_members.extend(Member(match, role, name, row=r) for role, name in names)

    return SimTab(code, code, roles, fr, final_columns, match_col, dr, draft_cols, draft_match_col,
                  final_members, absent, final_prefilled, draft_members, draft_prefilled)


# ---------------------------------------------------------------- pair history

@dataclass
class PairHistory:
    ids_by_student: dict[str, set[str]]        # class-name key -> pair IDs (never "_0")
    names_by_key: dict[str, str]
    absences_by_student: dict[str, set[str]]   # class-name key -> simulation codes they missed

    def shared_ids(self, a_key: str, b_key: str) -> list[str]:
        return sorted(self.ids_by_student.get(a_key, set()) & self.ids_by_student.get(b_key, set()))

    def linked_to_class(self, students: list[Student]) -> tuple["PairHistory", list[str]]:
        """Re-key the history by class-list student, so "Zach" on an old tab counts for "Zachary".

        Returns the re-keyed history and the tab names that could not be linked with confidence (reported to
        the instructor, never guessed). A name already matching the class list exactly is kept as is.
        """
        ids: dict[str, set[str]] = {}
        absences: dict[str, set[str]] = {}
        names: dict[str, str] = {s.key: s.name for s in students}
        unlinked = []
        for key in set(self.ids_by_student) | set(self.absences_by_student):
            res = resolve_class_name(self.names_by_key[key], students)
            if not res.confident:
                unlinked.append(self.names_by_key[key] + (f" (maybe {res.student.name})" if res.student else ""))
                continue
            ids.setdefault(res.student.key, set()).update(self.ids_by_student.get(key, set()))
            absences.setdefault(res.student.key, set()).update(self.absences_by_student.get(key, set()))
        return PairHistory(ids, names, absences), sorted(unlinked)


def read_pair_history(wb, exclude_code: str | None = None) -> PairHistory:
    """Collect every pair ID on every tab, attached to the student name to its left on the same row.

    Later tabs repeat earlier IDs next to each name, so the same ID may be seen several times; that is harmless.
    IDs for `exclude_code` (the simulation being worked on) are skipped.
    """
    ids: dict[str, set[str]] = {}
    absences: dict[str, set[str]] = {}
    names: dict[str, str] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            current_name = None
            for c in row:
                v = c.value
                if looks_like_class_name(v):
                    current_name = clean(v)
                    continue
                m = PAIR_ID.match(clean(v)) if isinstance(v, str) else None
                if not m or not current_name or m.group(1) == exclude_code:
                    continue
                key = class_name_key(current_name)
                names.setdefault(key, current_name)
                if m.group(2) == "0":
                    absences.setdefault(key, set()).add(m.group(1))
                else:
                    ids.setdefault(key, set()).add(clean(v))
    return PairHistory(ids, names, absences)
